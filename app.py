"""ERP Pro - Factory ERP (Flask + SQLite)
Run:  pip install flask  ->  python app.py  ->  http://localhost:5000
Default logins:  admin / admin123   |   hr / hr123   (change these passwords after the first login)
"""
import sqlite3, os, io, uuid
from datetime import date, timedelta
import calendar, re
from markupsafe import Markup
from functools import wraps
from flask import Flask, request, redirect, session, g, render_template_string, flash, url_for, send_file
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 40 * 1024 * 1024
UPLOADS = os.path.join(os.path.dirname(os.path.abspath(__file__)), "uploads"); os.makedirs(UPLOADS, exist_ok=True)
app.secret_key = os.environ.get("ERP_SECRET", "change-this-secret-key")
DB = os.path.join(os.path.dirname(__file__), "erp.db")

SCHEMA = """
CREATE TABLE IF NOT EXISTS users(id INTEGER PRIMARY KEY, username TEXT UNIQUE, password TEXT, role TEXT);
CREATE TABLE IF NOT EXISTS employees(id INTEGER PRIMARY KEY, name TEXT, cnic TEXT, phone TEXT,
  department TEXT, designation TEXT, salary REAL, join_date TEXT, active INTEGER DEFAULT 1);
CREATE TABLE IF NOT EXISTS attendance(id INTEGER PRIMARY KEY, emp_id INTEGER, day TEXT, status TEXT,
  UNIQUE(emp_id, day));
CREATE TABLE IF NOT EXISTS expenses(id INTEGER PRIMARY KEY, day TEXT, category TEXT, description TEXT, amount REAL);
CREATE TABLE IF NOT EXISTS inventory(id INTEGER PRIMARY KEY, item TEXT, sku TEXT, qty REAL, unit TEXT,
  min_qty REAL, price REAL);
CREATE TABLE IF NOT EXISTS leaves(id INTEGER PRIMARY KEY, emp_id INTEGER, start_date TEXT, end_date TEXT, reason TEXT, status TEXT);
CREATE TABLE IF NOT EXISTS documents(id INTEGER PRIMARY KEY, emp_id INTEGER, kind TEXT, stored TEXT, orig TEXT, created TEXT);
CREATE TABLE IF NOT EXISTS salary_history(id INTEGER PRIMARY KEY, emp_id INTEGER, day TEXT, basic REAL, note TEXT);
CREATE TABLE IF NOT EXISTS holidays(id INTEGER PRIMARY KEY, day TEXT, name TEXT);
CREATE TABLE IF NOT EXISTS corrections(id INTEGER PRIMARY KEY, emp_id INTEGER, day TEXT, new_status TEXT, check_in TEXT, check_out TEXT, reason TEXT, state TEXT, created TEXT);
CREATE TABLE IF NOT EXISTS overtime(id INTEGER PRIMARY KEY, emp_id INTEGER, day TEXT, hours REAL);
CREATE TABLE IF NOT EXISTS loans(id INTEGER PRIMARY KEY, emp_id INTEGER, amount REAL, installment REAL, start_month TEXT, note TEXT);
"""

DATABASE_URL = os.environ.get("DATABASE_URL")  # set this to use PostgreSQL, otherwise SQLite is used

class PG:
    """PostgreSQL wrapper that offers the same interface as SQLite."""
    def __init__(s):
        import psycopg2, psycopg2.extras
        s.c = psycopg2.connect(DATABASE_URL); s.x = psycopg2.extras.DictCursor
    def execute(s, q, p=()):
        cur = s.c.cursor(cursor_factory=s.x); cur.execute(q.replace("?", "%s"), p); return cur
    def commit(s): s.c.commit()
    def rollback(s): s.c.rollback()
    def close(s): s.c.close()

def connect():
    if DATABASE_URL: return PG()
    c = sqlite3.connect(DB); c.row_factory = sqlite3.Row; return c

def db():
    if "db" not in g: g.db = connect()
    return g.db

def insert(q, p):
    """INSERT and return the new row id (works on SQLite and PostgreSQL)."""
    if DATABASE_URL: return db().execute(q + " RETURNING id", p).fetchone()[0]
    return db().execute(q, p).lastrowid

@app.teardown_appcontext
def close(_):
    d = g.pop("db", None)
    if d: d.close()

def purge_employee(c, where, args=()):
    """Permanently remove employee(s) with their attendance, leaves, overtime, loans and login, then reset the ID counter."""
    c.execute(f"DELETE FROM attendance WHERE emp_id IN (SELECT id FROM employees WHERE {where})", args)
    for r in c.execute(f"SELECT stored FROM documents WHERE emp_id IN (SELECT id FROM employees WHERE {where})", args).fetchall():
        try: os.remove(os.path.join(UPLOADS, r[0]))
        except OSError: pass
    for t in ("leaves", "overtime", "loans", "documents", "salary_history", "corrections"):
        c.execute(f"DELETE FROM {t} WHERE emp_id IN (SELECT id FROM employees WHERE {where})", args)
    c.execute(f"DELETE FROM users WHERE emp_id IN (SELECT id FROM employees WHERE {where})", args)
    c.execute(f"DELETE FROM employees WHERE {where}", args)
    if DATABASE_URL:  # reset the ID counter on PostgreSQL
        c.execute("SELECT setval(pg_get_serial_sequence('employees','id'), COALESCE((SELECT MAX(id) FROM employees),0)+1, false)")

def init():
    c = connect()
    for st in SCHEMA.split(";"):
        if st.strip():
            if DATABASE_URL: st = st.replace("INTEGER PRIMARY KEY", "SERIAL PRIMARY KEY").replace("REAL", "DOUBLE PRECISION")
            c.execute(st)
    c.commit()
    TXT = "staff_code marital_status blood_group whatsapp permanent_address city province postal_code country skills prev_employment references_txt branch manager emp_type probation_start probation_end confirmation_date shift work_location status pay_method bank_name account_title account_no salary_freq emergency_rel emergency_alt emergency_address father_name gender dob passport nationality email address emergency_name emergency_phone".split()
    NUM = "allowances house_rent medical_allowance bonus ot_rate tax_deduction".split()
    for tbl, col, typ in ([("users", "emp_id", "INTEGER"), ("attendance", "check_in", "TEXT"), ("attendance", "check_out", "TEXT"), ("attendance", "late_min", "INTEGER"),
                           ("attendance", "early_min", "INTEGER"), ("leaves", "leave_type", "TEXT")]
                          + [("employees", k, "TEXT") for k in TXT] + [("employees", k, "REAL") for k in NUM]):
        try: c.execute(f"ALTER TABLE {tbl} ADD COLUMN {col} {typ}"); c.commit()
        except Exception: c.rollback()
    for r in c.execute("SELECT id FROM employees WHERE staff_code IS NULL OR staff_code=''").fetchall():   # automatic staff codes for existing employees
        c.execute("UPDATE employees SET staff_code=? WHERE id=?", (f"EMP-{r[0]:04d}", r[0]))
    c.commit()
    purge_employee(c, "active=0"); c.commit()   # clean old "inactive" employees so IDs restart from 1
    if not c.execute("SELECT 1 FROM users").fetchone():
        c.execute("INSERT INTO users(username,password,role) VALUES(?,?,?)", ("admin", generate_password_hash("admin123"), "admin"))
        c.execute("INSERT INTO users(username,password,role) VALUES(?,?,?)", ("hr", generate_password_hash("hr123"), "hr"))
    c.commit(); c.close()

def login_required(*roles):
    def deco(f):
        @wraps(f)
        def w(*a, **k):
            if "user" not in session: return redirect(url_for("login"))
            if roles and session["role"] not in roles:
                return "Access denied: you do not have permission to open this page.", 403
            return f(*a, **k)
        return w
    return deco

# ---------- UI ----------
BASE = """<!doctype html><html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1"><title>ERP Pro</title>
<link rel="preconnect" href="https://fonts.googleapis.com"><link href="https://fonts.googleapis.com/css2?family=Plus+Jakarta+Sans:wght@400;500;600;700;800&display=swap" rel="stylesheet">
<style>
:root{--bg:#f2f4f9;--ink:#0f1b2d;--mut:#6b778c;--navy:#0b1628;--navy2:#13233d;--acc:#f5a524;--acc2:#e0900f;--card:#fff;--line:#e4e8f0;--blue:#2f5fe3;--red:#c0362c;--shadow:0 1px 2px rgba(15,27,45,.04),0 8px 24px rgba(15,27,45,.06)}
*{box-sizing:border-box}
body{margin:0;font-family:'Plus Jakarta Sans',Segoe UI,system-ui,sans-serif;background:var(--bg);color:var(--ink);display:flex;min-height:100vh;font-size:14.5px;-webkit-font-smoothing:antialiased}
.ic{width:19px;height:19px;flex-shrink:0}
.side{width:252px;background:linear-gradient(180deg,var(--navy),var(--navy2));color:#aab6c9;padding:20px 14px;position:sticky;top:0;height:100vh;overflow-y:auto;flex-shrink:0;display:flex;flex-direction:column}
.brand{display:flex;gap:12px;align-items:center;padding:4px 8px 22px}
.logo{width:40px;height:40px;border-radius:12px;background:linear-gradient(135deg,#ffc857,var(--acc2));color:var(--navy);display:grid;place-items:center;font-weight:800;font-size:20px;box-shadow:0 6px 16px rgba(245,165,36,.35)}
.brand b{display:block;color:#fff;font-size:18px;letter-spacing:.2px}.brand small{color:#7f8da3;font-size:11.5px;letter-spacing:.6px;text-transform:uppercase}
.grp{font-size:10.5px;letter-spacing:1.2px;text-transform:uppercase;color:#5f6e86;margin:18px 10px 6px;font-weight:700}
.side a{display:flex;align-items:center;gap:12px;padding:10px 12px;margin:2px 0;border-radius:10px;color:#b7c2d4;text-decoration:none;font-weight:500;transition:.15s}
.side a:hover{background:rgba(255,255,255,.07);color:#fff}
.side a.on{background:linear-gradient(90deg,rgba(245,165,36,.2),rgba(245,165,36,.05));color:#fff;box-shadow:inset 3px 0 0 var(--acc)}
.side a.on .ic{color:var(--acc)}
.foot{margin-top:auto;padding-top:14px;border-top:1px solid rgba(255,255,255,.08)}
.wrap{flex:1;min-width:0;display:flex;flex-direction:column}
.top{display:flex;justify-content:space-between;align-items:center;padding:14px 32px;background:rgba(255,255,255,.85);backdrop-filter:blur(8px);border-bottom:1px solid var(--line);position:sticky;top:0;z-index:5}
.crumb{color:var(--mut)}.crumb i{margin:0 6px;font-style:normal;opacity:.5}.crumb b{color:var(--ink)}
.who{display:flex;align-items:center;gap:12px}.who .date{color:var(--mut);font-size:13px;margin-right:6px}
.who small{display:block;color:var(--mut);font-size:11.5px}.who b{font-size:13.5px}
.av{width:38px;height:38px;border-radius:50%;background:linear-gradient(135deg,var(--navy2),#2b4469);color:#fff;display:grid;place-items:center;font-weight:700}
main{padding:30px 32px 48px;max-width:1500px;width:100%}
main h2{font-size:26px;font-weight:800;margin:0 0 20px;letter-spacing:-.3px}
main h3{font-size:16px;margin:0 0 12px}
.card{background:var(--card);border:1px solid var(--line);border-radius:16px;padding:22px;margin-bottom:22px;box-shadow:var(--shadow);overflow-x:auto}
table{width:100%;border-collapse:collapse}
th{text-align:left;padding:11px 12px;font-size:11.5px;text-transform:uppercase;letter-spacing:.7px;color:var(--mut);background:#f7f9fc;border-bottom:1px solid var(--line);font-weight:700;white-space:nowrap}
th:first-child{border-radius:10px 0 0 10px}th:last-child{border-radius:0 10px 10px 0}
td{padding:12px;border-bottom:1px solid #eef1f6;vertical-align:middle}
tr:hover td{background:#fafbfe}tr:last-child td{border-bottom:0}
td a{color:var(--blue);font-weight:600;text-decoration:none}td a+form{margin-left:8px}td a:hover{text-decoration:underline}
input,select{padding:10px 13px;border:1px solid #d5dbe6;border-radius:10px;font:inherit;background:#fff;color:var(--ink);min-width:0;transition:.15s}
input:focus,select:focus{outline:0;border-color:var(--acc);box-shadow:0 0 0 4px rgba(245,165,36,.18)}
button{padding:10px 18px;border:0;border-radius:10px;font:inherit;font-weight:700;background:var(--acc);color:var(--navy);cursor:pointer;transition:.15s;box-shadow:0 2px 6px rgba(245,165,36,.3)}
button:hover{background:var(--acc2);transform:translateY(-1px)}
td button{padding:6px 12px;font-size:12.5px;box-shadow:none}
form[action*="/delete/"] button{background:#fdecea;color:var(--red)}form[action*="/delete/"] button:hover{background:var(--red);color:#fff}
form[action*="/reject"] button{background:#fdecea;color:var(--red)}form[action*="/approve"] button{background:#e3f6ec;color:#12875b}
form.row{display:flex;flex-wrap:wrap;gap:10px;margin-bottom:16px;align-items:center}
.low{color:var(--red);font-weight:700}
.flash{background:#fff8e6;border:1px solid #fde3a7;border-left:4px solid var(--acc);padding:12px 16px;border-radius:10px;margin-bottom:18px;font-weight:500}
.hero{margin-bottom:24px}.hero h2{margin-bottom:4px}.hero p{margin:0;color:var(--mut)}
.stats{display:grid;grid-template-columns:repeat(3,1fr);gap:18px;margin-bottom:24px}@media(max-width:1100px){.stats{grid-template-columns:repeat(auto-fit,minmax(230px,1fr))}}
.sc{background:var(--card);border:1px solid var(--line);border-radius:16px;padding:20px;display:flex;gap:16px;align-items:flex-start;box-shadow:var(--shadow);text-decoration:none;color:inherit;transition:.2s}
a.sc:hover{transform:translateY(-3px);box-shadow:0 12px 30px rgba(15,27,45,.1)}
.sc>div{flex:1}.sc small{color:var(--mut);font-weight:600;display:block;margin-bottom:4px}.sc strong{font-size:27px;font-weight:800;letter-spacing:-.5px}.sc em{font-style:normal;font-size:14px;color:var(--mut);font-weight:600}
.ico{width:46px;height:46px;border-radius:13px;display:grid;place-items:center;flex-shrink:0}.ico .ic{width:22px;height:22px}
.ico.b{background:#e8efff;color:var(--blue)}.ico.g{background:#e3f6ec;color:#12875b}.ico.a{background:#fff1d6;color:#b9770e}.ico.p{background:#efe8ff;color:#6d3fd1}.ico.r{background:#fde8e8;color:var(--red)}
.bar{height:6px;background:#e9edf4;border-radius:9px;margin-top:10px;overflow:hidden}.bar i{display:block;height:100%;background:linear-gradient(90deg,#1fb37a,#12875b);border-radius:9px}
.qa{display:flex;flex-wrap:wrap;gap:12px}.qa a{display:flex;gap:10px;align-items:center;padding:12px 18px;border-radius:12px;background:#f5f7fc;border:1px solid var(--line);color:var(--ink);text-decoration:none;font-weight:600;transition:.15s}
.qa a:hover{background:var(--navy);color:#fff;border-color:var(--navy)}
.fgrid{display:grid;grid-template-columns:repeat(auto-fit,minmax(250px,1fr));gap:16px 20px}
.fld label{display:block;font-size:12.5px;font-weight:700;color:var(--mut);margin-bottom:6px}.fld input,.fld select{width:100%}
.sec{grid-column:1/-1;font-size:12px;text-transform:uppercase;letter-spacing:1px;color:var(--navy2);font-weight:800;margin-top:10px;padding-top:16px;border-top:1px solid var(--line)}.sec:first-child{border-top:0;margin-top:0;padding-top:0}
.ph{display:flex;flex-direction:column;gap:8px}.fld form a,form a.cx{color:var(--mut);font-weight:600;margin-left:12px;text-decoration:none}
.tb{display:flex;flex-wrap:wrap;gap:12px;align-items:center;margin-bottom:16px}.tb .s{position:relative;flex:1;min-width:240px}
.tb .s .ic{position:absolute;left:13px;top:50%;transform:translateY(-50%);color:#8794ab;width:18px;height:18px}.tb .s input{width:100%;padding-left:42px}
.tb .cnt{color:var(--mut);font-size:13px;font-weight:600;margin-left:auto}.empty{text-align:center;color:var(--mut);padding:28px;display:none}
.phead a,.pg a{color:var(--blue);font-weight:600;text-decoration:none}.phead a:hover,.pg a:hover{text-decoration:underline}td.num,td.idc,td.nw,td.noprint{white-space:nowrap}td.noprint form{margin-left:6px}.av2{width:38px;height:38px;border-radius:50%;background:#e8efff;color:var(--blue);display:grid;place-items:center;font-weight:800;margin-right:12px;flex-shrink:0}
.nm{display:flex;align-items:center}.nm small,td small{display:block;color:var(--mut);font-size:12px}
.phead{display:flex;justify-content:space-between;align-items:center;margin-bottom:20px;gap:12px;flex-wrap:wrap}.phead h2{margin:0}
.phead button{display:inline-flex;gap:8px;align-items:center}
.pg{display:grid;grid-template-columns:repeat(auto-fit,minmax(230px,1fr));gap:22px 18px}.pg small{display:block;color:var(--mut);font-size:11.5px;text-transform:uppercase;letter-spacing:.7px;font-weight:700;margin-bottom:4px}
.pgh{display:flex;gap:20px;align-items:center;margin-bottom:26px}.pgh .av2{width:78px;height:78px;font-size:32px;margin:0}.pgh h2{margin:0 0 4px}.pgh p{margin:0;color:var(--mut)}
input[type=checkbox]{width:17px;height:17px;min-width:17px;padding:0;accent-color:var(--acc);cursor:pointer}
.bulk{position:fixed;left:50%;bottom:22px;transform:translateX(-50%);display:none;gap:12px;align-items:center;flex-wrap:wrap;justify-content:center;max-width:94vw;background:var(--navy);color:#fff;padding:12px 18px;border-radius:16px;box-shadow:0 18px 50px rgba(11,22,40,.45);z-index:40}
.bulk.show{display:flex}.bulk select,.bulk input[type=month]{padding:8px 10px;border-radius:9px;border:0;font-size:13.5px}.bulk .o{display:flex;gap:10px;align-items:center;font-size:13px}
.bulk label{display:flex;gap:6px;align-items:center;font-size:13px}.bulk button{padding:9px 16px;display:inline-flex;gap:8px;align-items:center}
.bulk .danger{background:#fdecea;color:var(--red);box-shadow:none}.bulk .danger:hover{background:var(--red);color:#fff}.bulk .ghost{background:rgba(255,255,255,.12);color:#fff;box-shadow:none}
table.ft tr.sel td{background:#fff8e6}table.att tr.sel td.nm{background:#fff4d6!important}
textarea{width:100%;padding:10px 13px;border:1px solid #d5dbe6;border-radius:10px;font:inherit;resize:vertical;background:#fff;color:var(--ink)}textarea:focus{outline:0;border-color:var(--acc);box-shadow:0 0 0 4px rgba(245,165,36,.18)}
.fld.wide{grid-column:1/-1}.tabs{display:flex;gap:4px;flex-wrap:wrap;border-bottom:1px solid var(--line);margin-bottom:22px}
.tab{background:none;color:var(--mut);box-shadow:none;border-radius:10px 10px 0 0;padding:11px 16px;font-weight:700;border-bottom:3px solid transparent}.tab:hover{background:#f5f7fc;transform:none;color:var(--ink)}
.tab.on{color:var(--ink);border-bottom-color:var(--acc);background:#fff8e6}.panel{display:none}.panel.on{display:block}
.av2{overflow:hidden}.av2 img{width:100%;height:100%;object-fit:cover}
.bd{display:inline-block;padding:3px 11px;border-radius:99px;font-size:12px;font-weight:700}.b-Active{background:#e3f6ec;color:#12875b}.b-OnLeave{background:#e8efff;color:#2f5fe3}.b-Resigned{background:#fff1d6;color:#b9770e}.b-Terminated{background:#fde8e8;color:#c0362c}
.doc{display:flex;gap:10px;align-items:center;flex-wrap:wrap;margin:6px 0}.doc form{display:inline}.thumb{height:84px;border-radius:8px;border:1px solid var(--line)}
@media print{.side,.top,.noprint{display:none!important}main{padding:0}.card{box-shadow:none;border:0}}
@media(max-width:900px){body{flex-direction:column}.side{width:100%;height:auto;position:static;flex-direction:row;overflow-x:auto;padding:10px;align-items:center}
.brand,.grp,.foot{display:none}.side a{white-space:nowrap}.top{padding:12px 16px}.who .date{display:none}main{padding:20px 16px}}
</style></head><body>
{% if session.user %}<aside class="side"><div class="brand"><span class="logo">E</span><div><b>ERP Pro</b><small>Next Gen</small></div></div>
{% for t, items in nav %}<div class="grp">{{t}}</div>{% for h,l,i,r in items %}<a href="{{h}}" class="{{'on' if h==active}}">{{icon(i)}}<span>{{l}}</span></a>{% endfor %}{% endfor %}
<div class="foot"><a href="/logout">{{icon('logout')}}<span>Logout</span></a></div></aside>{% endif %}
<div class="wrap">{% if session.user %}<header class="top"><div class="crumb">ERP Pro<i>/</i><b>{{crumb}}</b></div>
<div class="who"><span class="date">{{today}}</span><span class="av">{{initial}}</span><div><b>{{session.user}}</b><small>{{session.role|capitalize}}</small></div></div></header>{% endif %}
<main>{% for m in get_flashed_messages() %}<div class="flash">{{m}}</div>{% endfor %}{{body|safe}}</main></div>
<script>
(function(){
var IC='{{icon("search")}}';
function codeOf(t){var m=t.match(/[(][+]([0-9]+)[)]/);if(m)return m[1];m=t.trim().match(/^[+]?([0-9]{1,4})$/);return m?m[1]:'';}
document.querySelectorAll('.ph').forEach(function(w){
 var cc=w.querySelector('.cc'),pn=w.querySelector('.pn'),h=w.querySelector('input[type=hidden]');
 function sync(){var c=codeOf(cc.value),n=pn.value.replace(/[^0-9]/g,'');if(c)n=n.replace(/^0+/,'');h.value=n?(c?'+'+c+' '+n:n):'';}
 var v=h.value.trim(),m=v.match(/^[+]([0-9]+) *(.*)$/);
 if(m){var opts=[].map.call(document.querySelectorAll('#countries option'),function(o){return o.value;});
  cc.value=opts.filter(function(x){return codeOf(x)===m[1];})[0]||('+'+m[1]);pn.value=m[2].replace(/[^0-9]/g,'');}
 else{cc.value='Pakistan (+92)';pn.value=v.replace(/[^0-9]/g,'').replace(/^0+/,'');}
 cc.addEventListener('focus',function(){cc.select();});cc.addEventListener('input',sync);cc.addEventListener('change',sync);
 pn.addEventListener('input',function(){pn.value=pn.value.replace(/[^0-9]/g,'');sync();});sync();
});
document.querySelectorAll('input.cnic').forEach(function(i){
 function f(){var d=i.value.replace(/[^0-9]/g,'').slice(0,13),o=d.slice(0,5);if(d.length>5)o+='-'+d.slice(5,12);if(d.length>12)o+='-'+d.slice(12);i.value=o;}
 i.addEventListener('input',f);f();
});
document.querySelectorAll('.tabs').forEach(function(tb){var f=tb.closest('form');
 function show(id){tb.querySelectorAll('.tab').forEach(function(b){b.classList.toggle('on',b.getAttribute('data-t')===id);});f.querySelectorAll('.panel').forEach(function(p){p.classList.toggle('on',p.id==='p-'+id);});}
 tb.addEventListener('click',function(e){var b=e.target.closest('.tab');if(b)show(b.getAttribute('data-t'));});
 f.addEventListener('invalid',function(e){var p=e.target.closest('.panel');if(p)show(p.id.slice(2));},true);
});
document.querySelectorAll('table.ft').forEach(function(t){
 var rows=[].slice.call(t.rows).slice(1);if(!rows.length)return;
 var specs=[[/^dept|^department/i,'All departments'],[/^status$/i,'All statuses']],hs=[].slice.call(t.rows[0].cells),fl=[];
 specs.forEach(function(sp){var ix=hs.findIndex(function(h){return sp[0].test(h.textContent.trim());});if(ix>=0)fl.push({ix:ix,label:sp[1]});});
 var tb=document.createElement('div');tb.className='tb';
 tb.innerHTML='<div class="s">'+IC+'<input type="search" placeholder="'+(t.getAttribute('data-ph')||'Search...')+'"></div>'+fl.map(function(f){return '<select><option value="">'+f.label+'</option></select>';}).join('')+'<span class="cnt"></span>';
 t.parentNode.insertBefore(tb,t);
 var q=tb.querySelector('input'),sels=[].slice.call(tb.querySelectorAll('select')),cnt=tb.querySelector('.cnt'),em=document.createElement('div');
 em.className='empty';em.textContent='No results match your search.';t.parentNode.insertBefore(em,t.nextSibling);
 fl.forEach(function(f,i){var seen={};rows.forEach(function(r){var v=r.cells[f.ix].textContent.trim();if(v)seen[v]=1;});
  Object.keys(seen).sort().forEach(function(v){var o=document.createElement('option');o.value=o.textContent=v;sels[i].appendChild(o);});sels[i].addEventListener('change',run);});
 function txt(r){return [].map.call(r.cells,function(c){var s=c.querySelector('select');return s?s.value:c.textContent;}).join(' ').toLowerCase();}
 function run(){
  var s=q.value.trim().toLowerCase(),text=/[^0-9 +().-]/.test(s),d=s.replace(/[^0-9]/g,''),d0=d.replace(/^0+/,''),n=0;
  rows.forEach(function(r){
   var ok=fl.every(function(f,i){return !sels[i].value||r.cells[f.ix].textContent.trim()===sels[i].value;});
   if(ok&&s){
    var nc=r.querySelectorAll('td.num'),ids=r.querySelectorAll('td.idc');
    if(!text&&d&&(nc.length||ids.length)){
     ok=[].some.call(nc,function(c){var x=c.textContent.replace(/[^0-9]/g,'');return x.indexOf(d)>-1||(d0&&x.indexOf(d0)>-1);})||[].some.call(ids,function(c){return c.textContent.trim()===d;});
    }else ok=txt(r).indexOf(s)>-1;
   }
   r.style.display=ok?'':'none';if(ok)n++;
  });
  cnt.textContent='Showing '+n+' of '+rows.length;em.style.display=n?'none':'block';
 }
 q.addEventListener('input',run);q.addEventListener('keydown',function(e){if(e.key==='Enter')e.preventDefault();});run();
});
var bk=document.getElementById('bulk'),bt=document.querySelector('table.bt');
if(bk&&bt){
 var bn=document.getElementById('bn'),kind=document.getElementById('bkind'),bm=document.getElementById('bm'),bf=document.getElementById('bf'),bd=document.getElementById('bd'),opt=document.getElementById('bopt');
 function sel(){return [].filter.call(bt.querySelectorAll('input.rc:checked'),function(c){return c.closest('tr').style.display!=='none';}).map(function(c){return c.value;});}
 function upd(){var s=sel();bk.className='bulk'+(s.length?' show':'');bn.textContent=s.length;[].forEach.call(bt.querySelectorAll('input.rc'),function(c){c.closest('tr').classList.toggle('sel',c.checked);});opt.style.display=kind.value==='att'?'flex':'none';}
 bt.addEventListener('change',function(e){var t=e.target;if(t.classList.contains('all'))[].forEach.call(bt.querySelectorAll('input.rc'),function(c){if(c.closest('tr').style.display!=='none')c.checked=t.checked;});if(t.classList.contains('all')||t.classList.contains('rc'))upd();});
 ['input','change'].forEach(function(ev){document.addEventListener(ev,function(e){if(e.target.closest&&e.target.closest('.tb'))setTimeout(upd,0);});});
 kind.addEventListener('change',upd);
 document.getElementById('bdl').addEventListener('click',function(){var s=sel();if(!s.length)return;location.href='/bulk/download?ids='+s.join(',')+'&kind='+kind.value+'&fmt='+bf.value+'&month='+bm.value+'&detail='+(bd.checked?1:0);});
 document.getElementById('bclr').addEventListener('click',function(){[].forEach.call(document.querySelectorAll('input.rc,input.all'),function(c){c.checked=false;});upd();});
 document.getElementById('bdel').addEventListener('click',function(){var s=sel();if(!s.length||!confirm('Permanently delete '+s.length+' employee(s)? Their attendance, leaves, overtime, loans and logins will be deleted too.'))return;
  var f=document.createElement('form');f.method='post';f.action='/employees/bulk_delete';f.innerHTML='<input type="hidden" name="ids" value="'+s.join(',')+'"><input type="hidden" name="next" value="'+location.pathname+'">';document.body.appendChild(f);f.submit();});
 upd();
}
})();
</script></body></html>"""

ICONS = {
 "home": '<path d="M3 11l9-8 9 8v10a1 1 0 0 1-1 1h-5v-7H9v7H4a1 1 0 0 1-1-1z"/>',
 "users": '<path d="M17 21v-2a4 4 0 0 0-4-4H5a4 4 0 0 0-4 4v2"/><circle cx="9" cy="7" r="4"/><path d="M23 21v-2a4 4 0 0 0-3-3.87M16 3.13a4 4 0 0 1 0 7.75"/>',
 "user": '<path d="M20 21v-2a4 4 0 0 0-4-4H8a4 4 0 0 0-4 4v2"/><circle cx="12" cy="7" r="4"/>',
 "calendar": '<rect x="3" y="4" width="18" height="18" rx="2"/><path d="M16 2v4M8 2v4M3 10h18"/>',
 "chart": '<path d="M18 20V10M12 20V4M6 20v-6"/>',
 "file": '<path d="M14 2H6a2 2 0 0 0-2 2v16a2 2 0 0 0 2 2h12a2 2 0 0 0 2-2V8z"/><path d="M14 2v6h6M16 13H8M16 17H8"/>',
 "wallet": '<rect x="2" y="5" width="20" height="14" rx="2"/><path d="M2 10h20"/>',
 "clock": '<circle cx="12" cy="12" r="10"/><path d="M12 6v6l4 2"/>',
 "dollar": '<path d="M12 1v22M17 5H9.5a3.5 3.5 0 0 0 0 7h5a3.5 3.5 0 0 1 0 7H6"/>',
 "receipt": '<path d="M4 2v20l3-2 3 2 3-2 3 2 3-2V2z"/><path d="M8 8h8M8 12h8"/>',
 "box": '<path d="M21 8l-9-5-9 5v8l9 5 9-5z"/><path d="M3 8l9 5 9-5M12 13v8"/>',
 "shield": '<path d="M12 22s8-4 8-10V5l-8-3-8 3v7c0 6 8 10 8 10z"/>',
 "download": '<path d="M21 15v4a2 2 0 0 1-2 2H5a2 2 0 0 1-2-2v-4M7 10l5 5 5-5M12 15V3"/>',
 "key": '<circle cx="7.5" cy="15.5" r="4.5"/><path d="M21 2l-9.6 9.6M15.5 7.5l3 3"/>',
 "logout": '<path d="M9 21H5a2 2 0 0 1-2-2V5a2 2 0 0 1 2-2h4M16 17l5-5-5-5M21 12H9"/>',
 "check": '<path d="M22 11.08V12a10 10 0 1 1-5.93-9.14"/><path d="M22 4L12 14.01l-3-3"/>',
 "alert": '<path d="M10.29 3.86L1.82 18a2 2 0 0 0 1.71 3h16.94a2 2 0 0 0 1.71-3L13.71 3.86a2 2 0 0 0-3.42 0z"/><path d="M12 9v4M12 17h.01"/>',
 "lock": '<rect x="3" y="11" width="18" height="11" rx="2"/><path d="M7 11V7a5 5 0 0 1 10 0v4"/>',
 "phone": '<path d="M22 16.92v3a2 2 0 0 1-2.18 2 19.79 19.79 0 0 1-8.63-3.07 19.5 19.5 0 0 1-6-6 19.79 19.79 0 0 1-3.07-8.67A2 2 0 0 1 4.11 2h3a2 2 0 0 1 2 1.72c.13.96.36 1.9.7 2.81a2 2 0 0 1-.45 2.11L8.09 9.91a16 16 0 0 0 6 6l1.27-1.27a2 2 0 0 1 2.11-.45c.91.34 1.85.57 2.81.7A2 2 0 0 1 22 16.92z"/>',
 "mail": '<rect x="2" y="4" width="20" height="16" rx="2"/><path d="m22 7-10 5L2 7"/>',
 "search": '<circle cx="11" cy="11" r="8"/><path d="M21 21l-4.35-4.35"/>',
 "plus": '<path d="M12 5v14M5 12h14"/>',
}
app.jinja_env.globals["icon"] = lambda n, cls="ic": Markup(
    f'<svg class="{cls}" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="1.8" stroke-linecap="round" stroke-linejoin="round">{ICONS[n]}</svg>')

HRA, ALL = ("admin", "hr"), ("admin", "hr", "employee")
NAV = [("Overview", [("/", "Dashboard", "home", HRA), ("/me", "My portal", "user", ("employee",))]),
       ("People", [("/employees", "Employees", "users", HRA), ("/attendance", "Attendance", "calendar", HRA),
                   ("/leaves", "Leaves", "file", HRA),
                   ("/corrections", "Corrections", "check", HRA), ("/holidays", "Holidays", "calendar", HRA)]),
       ("Payroll", [("/payroll", "Payroll", "wallet", HRA), ("/overtime", "Overtime", "clock", HRA), ("/loans", "Advance / Loan", "dollar", HRA)]),
       ("Operations", [("/expenses", "Expenses", "receipt", ("admin",)), ("/inventory", "Inventory", "box", ("admin",))]),
       ("System", [("/exports", "Export to Excel", "download", HRA), ("/users", "Users", "shield", ("admin",)),
                   ("/password", "Change password", "key", ALL)])]

@app.context_processor
def inject_nav():
    role = session.get("role")
    if not role: return {}
    secs = [(t, [i for i in items if role in i[3]]) for t, items in NAV]; secs = [x for x in secs if x[1]]
    path = request.path; m = re.match(r"/edit/(\w+)/", path)
    if m and m.group(1) in CFG: path = CFG[m.group(1)]["back"]
    elif path.startswith("/employee/"): path = "/employees"
    elif path.startswith("/payslip"): path = "/me" if role == "employee" else "/payroll"
    hits = [i for _, it in secs for i in it if path == i[0] or (i[0] != "/" and path.startswith(i[0]))]
    best = max(hits, key=lambda i: len(i[0]), default=None)
    return dict(nav=secs, active=best[0] if best else None, crumb=best[1] if best else "ERP Pro",
                today=date.today().strftime("%A, %d %B %Y"), initial=session["user"][:1].upper())

def page(tpl, **ctx):
    from flask import render_template_string as r
    return r(BASE, body=r(tpl, **ctx))


LOGIN = """<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Sign in - ERP Pro</title>
<link rel="preconnect" href="https://fonts.googleapis.com"><link href="https://fonts.googleapis.com/css2?family=Plus+Jakarta+Sans:wght@400;500;600;700;800&display=swap" rel="stylesheet">
<style>
*{box-sizing:border-box}body{margin:0;font-family:'Plus Jakarta Sans',Segoe UI,system-ui,sans-serif;min-height:100vh;display:flex;background:#fff;color:#0f1b2d;-webkit-font-smoothing:antialiased}
.ic{width:20px;height:20px;flex-shrink:0}
.hero{flex:1.15;position:relative;overflow:hidden;background:linear-gradient(145deg,#0b1628 0%,#13233d 55%,#1b3358 100%);color:#fff;padding:56px 64px;display:flex;flex-direction:column;justify-content:space-between}
.hero:before{content:"";position:absolute;inset:0;background-image:linear-gradient(rgba(255,255,255,.045) 1px,transparent 1px),linear-gradient(90deg,rgba(255,255,255,.045) 1px,transparent 1px);background-size:44px 44px;mask-image:radial-gradient(ellipse at 30% 40%,#000 25%,transparent 75%)}
.hero:after{content:"";position:absolute;width:520px;height:520px;right:-170px;bottom:-190px;border-radius:50%;background:radial-gradient(circle,rgba(245,165,36,.55),rgba(245,165,36,0) 68%)}
.hero>*{position:relative;z-index:1}
.brand{display:flex;gap:14px;align-items:center}.logo{width:46px;height:46px;border-radius:14px;background:linear-gradient(135deg,#ffc857,#e0900f);color:#0b1628;display:grid;place-items:center;font-weight:800;font-size:23px;box-shadow:0 8px 22px rgba(245,165,36,.4)}
.brand b{font-size:22px;letter-spacing:.2px}.brand small{display:block;color:#8fa0bb;font-size:12px;letter-spacing:.8px;text-transform:uppercase}
h1{font-size:44px;line-height:1.1;margin:0 0 16px;font-weight:800;letter-spacing:-1px;max-width:520px}h1 span{color:#f5a524}
.lead{color:#a9b7cd;font-size:16.5px;line-height:1.6;max-width:480px;margin:0 0 34px}
.feat{display:grid;grid-template-columns:1fr 1fr;gap:14px;max-width:520px}
.feat div{display:flex;gap:12px;align-items:center;padding:13px 15px;border-radius:14px;background:rgba(255,255,255,.06);border:1px solid rgba(255,255,255,.09);backdrop-filter:blur(6px);font-weight:600;font-size:14px}
.feat .ic{color:#f5a524}.copy{color:#6f819f;font-size:13px}
.panel{flex:1;display:grid;place-items:center;padding:40px 24px;background:#f6f8fc}
.box{width:100%;max-width:410px;background:#fff;border:1px solid #e4e8f0;border-radius:22px;padding:42px 36px;box-shadow:0 24px 60px rgba(15,27,45,.1)}
.mlogo{display:none;margin-bottom:22px}.box h2{margin:0 0 6px;font-size:28px;font-weight:800;letter-spacing:-.5px}.sub{margin:0 0 28px;color:#6b778c}
label{display:block;font-weight:600;font-size:13px;margin-bottom:7px}.f{position:relative;margin-bottom:18px}
.f .ic{position:absolute;left:14px;top:50%;transform:translateY(-50%);color:#8794ab}
.f input{width:100%;padding:13px 14px 13px 44px;border:1px solid #d5dbe6;border-radius:12px;font:inherit;font-size:15px;transition:.15s;background:#fbfcfe}
.f input:focus{outline:0;border-color:#f5a524;background:#fff;box-shadow:0 0 0 4px rgba(245,165,36,.18)}
.eye{position:absolute;right:10px;top:50%;transform:translateY(-50%);background:none;border:0;color:#6b778c;font-weight:600;font-size:12.5px;cursor:pointer;padding:6px 8px;box-shadow:none}
.go{width:100%;padding:14px;border:0;border-radius:12px;font:inherit;font-weight:700;font-size:15.5px;color:#0b1628;background:linear-gradient(135deg,#ffc857,#f5a524);cursor:pointer;box-shadow:0 10px 24px rgba(245,165,36,.4);transition:.15s;margin-top:6px}
.go:hover{transform:translateY(-2px);box-shadow:0 14px 30px rgba(245,165,36,.5)}
.err{background:#fdecea;color:#b42318;border:1px solid #f8c9c4;border-radius:12px;padding:12px 14px;margin-bottom:20px;font-weight:600;font-size:14px}
.note{text-align:center;color:#8794ab;font-size:12.5px;margin-top:26px}
.stack{width:100%;max-width:410px}
.dev{display:flex;align-items:center;gap:18px;margin-top:20px;padding:20px 22px;background:#fff;border:1px solid #e4e8f0;border-radius:18px;box-shadow:0 10px 30px rgba(15,27,45,.07)}
.dev-img{width:92px;height:92px;border-radius:50%;object-fit:cover;border:3px solid #f5a524;box-shadow:0 6px 16px rgba(245,165,36,.35);flex-shrink:0}
.dev .dev-av{color:#fff;font-size:20px}.dev-av{display:grid;place-items:center;background:linear-gradient(135deg,#13233d,#2b4469);color:#fff;font-weight:800;font-size:20px;letter-spacing:1px}
.dev small{display:block;color:#8794ab;font-size:11px;text-transform:uppercase;letter-spacing:1px;font-weight:700}
.dev b{display:block;font-size:18px;margin:2px 0}.dev .role{color:#6b778c;font-size:13px;display:block;margin-bottom:8px}.dev a.ct{display:flex;align-items:center;gap:8px;color:#2f5fe3;font-size:13px;font-weight:600;text-decoration:none;margin-top:4px;word-break:break-all}.dev a.ct:hover{text-decoration:underline}.dev a.ct .ic{width:15px;height:15px;color:#f5a524}
@media(max-width:900px){.hero{display:none}.mlogo{display:flex}}
</style></head><body>
<section class="hero"><div class="brand"><span class="logo">E</span><div><b>ERP Pro</b><small>Next Gen</small></div></div>
<div><h1>Your entire factory, <span>one portal</span>.</h1><p class="lead">Employees, attendance, payroll, expenses and inventory in one place &mdash; clear and fast.</p>
<div class="feat"><div>{{icon('users')}}Employees</div><div>{{icon('calendar')}}Attendance</div><div>{{icon('wallet')}}Payroll &amp; slips</div><div>{{icon('box')}}Inventory</div></div></div>
<div class="copy">&copy; ERP Pro. All rights reserved. &middot; Developed by Mr Usman Qasim<br>0339-4023630 &middot; 2usmanqasim@gmail.com</div></section>
<section class="panel"><div class="stack"><div class="box"><div class="brand mlogo"><span class="logo">E</span><div><b>ERP Pro</b></div></div>
<h2>Welcome back</h2><p class="sub">Sign in to your account to continue.</p>
{% for m in get_flashed_messages() %}<div class="err">{{m}}</div>{% endfor %}
<form method="post"><label>Username</label><div class="f">{{icon('user')}}<input name="username" placeholder="Username" autocomplete="username" autofocus required></div>
<label>Password</label><div class="f">{{icon('lock')}}<input id="pw" type="password" name="password" placeholder="Password" autocomplete="current-password" required>
<button type="button" class="eye" onclick="var i=document.getElementById('pw');i.type=i.type==='password'?'text':'password'">Show</button></div>
<button class="go">Sign in</button></form><div class="note">Need help? Contact your Admin.</div></div>
<div class="dev">{% if photo %}<img class="dev-img" src="{{photo}}" alt="Mr Usman Qasim">{% else %}<span class="dev-img dev-av">UQ</span>{% endif %}
<div><small>Developed by</small><b>Mr Usman Qasim</b><span class="role">Software Developer</span><a class="ct" href="tel:03394023630">{{icon('phone')}}0339-4023630</a><a class="ct" href="mailto:2usmanqasim@gmail.com">{{icon('mail')}}2usmanqasim@gmail.com</a></div></div></div></section></body></html>"""

# ---------- Auth ----------
@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        u = db().execute("SELECT * FROM users WHERE username=?", (request.form["username"],)).fetchone()
        if u and check_password_hash(u["password"], request.form["password"]):
            session.update(user=u["username"], role=u["role"], emp_id=u["emp_id"]); return redirect("/me" if u["role"] == "employee" else "/")
        flash("Incorrect username or password.")
    photo = next((f"/static/developer.{e}" for e in ("jpg", "jpeg", "png", "webp")
                  if os.path.exists(os.path.join(app.static_folder, f"developer.{e}"))), None)
    return render_template_string(LOGIN, photo=photo)

@app.route("/logout")
def logout():
    session.clear(); return redirect("/login")

# ---------- Dashboard ----------
@app.route("/")
@login_required()
def dashboard():
    if session["role"] == "employee": return redirect("/me")
    d = db(); today = date.today().isoformat(); month = today[:7]
    s = dict(
        emps=d.execute("SELECT COUNT(*) FROM employees WHERE active=1 AND COALESCE(status,'Active') NOT IN ('Resigned','Terminated')").fetchone()[0],
        present=d.execute("SELECT COUNT(*) FROM attendance a JOIN employees e ON e.id=a.emp_id WHERE e.active=1 AND COALESCE(e.status,'Active') NOT IN ('Resigned','Terminated') AND a.day=? AND a.status='Present'", (today,)).fetchone()[0],
        exp=d.execute("SELECT COALESCE(SUM(amount),0) FROM expenses WHERE day LIKE ?", (month + "%",)).fetchone()[0],
        low=d.execute("SELECT COUNT(*) FROM inventory WHERE qty<=min_qty").fetchone()[0],
        payroll=d.execute("SELECT COALESCE(SUM(salary),0) FROM employees WHERE active=1 AND COALESCE(status,'Active') NOT IN ('Resigned','Terminated')").fetchone()[0],
        leaves=d.execute("SELECT COUNT(*) FROM leaves WHERE status='Pending'").fetchone()[0])
    rate = round(100 * s["present"] / s["emps"]) if s["emps"] else 0
    return page("""<div class="hero"><h2>Welcome back, {{session.user|capitalize}}</h2><p>Here is today's factory overview.</p></div>
    <div class="stats">
    <div class="sc"><span class="ico b">{{icon('users')}}</span><div><small>Employees</small><strong>{{s.emps}}</strong></div></div>
    <div class="sc"><span class="ico g">{{icon('check')}}</span><div><small>Present today</small><strong>{{s.present}} <em>/ {{s.emps}}</em></strong><div class="bar"><i style="width:{{rate}}%"></i></div></div></div>
    <a class="sc" href="/leaves"><span class="ico {{'r' if s.leaves else 'p'}}">{{icon('file')}}</span><div><small>Pending leaves</small><strong>{{s.leaves}}</strong></div></a>
    <div class="sc"><span class="ico a">{{icon('receipt')}}</span><div><small>Expenses (this month)</small><strong><em>Rs</em> {{'{:,.0f}'.format(s.exp)}}</strong></div></div>
    <div class="sc"><span class="ico p">{{icon('wallet')}}</span><div><small>Monthly payroll</small><strong><em>Rs</em> {{'{:,.0f}'.format(s.payroll)}}</strong></div></div>
    <div class="sc"><span class="ico {{'r' if s.low else 'g'}}">{{icon('alert')}}</span><div><small>Low stock items</small><strong>{{s.low}}</strong></div></div></div>
    <div class="card"><h3>Quick actions</h3><div class="qa">
    <a href="/attendance">{{icon('calendar')}}Mark attendance</a><a href="/employees">{{icon('plus')}}Add employee</a>
    <a href="/payroll">{{icon('wallet')}}Run payroll</a><a href="/leaves">{{icon('file')}}Leave requests</a>
    {% if session.role=='admin' %}<a href="/expenses">{{icon('receipt')}}Add expense</a><a href="/inventory">{{icon('box')}}Inventory</a>{% endif %}</div></div>""", s=s, rate=rate)

# ---------- Employees (admin + HR) ----------
COUNTRIES = [(x.rsplit(":", 1)[0], x.rsplit(":", 1)[1]) for x in "Afghanistan:93;Albania:355;Algeria:213;Andorra:376;Angola:244;Argentina:54;Armenia:374;Australia:61;Austria:43;Azerbaijan:994;Bahrain:973;Bangladesh:880;Belarus:375;Belgium:32;Benin:229;Bhutan:975;Bolivia:591;Bosnia and Herzegovina:387;Botswana:267;Brazil:55;Brunei:673;Bulgaria:359;Burkina Faso:226;Burundi:257;Cambodia:855;Cameroon:237;Canada:1;Chad:235;Chile:56;China:86;Colombia:57;DR Congo:243;Costa Rica:506;Croatia:385;Cuba:53;Cyprus:357;Czechia:420;Denmark:45;Djibouti:253;Ecuador:593;Egypt:20;El Salvador:503;Eritrea:291;Estonia:372;Ethiopia:251;Finland:358;France:33;Gabon:241;Gambia:220;Georgia:995;Germany:49;Ghana:233;Greece:30;Guatemala:502;Guinea:224;Haiti:509;Honduras:504;Hong Kong:852;Hungary:36;Iceland:354;India:91;Indonesia:62;Iran:98;Iraq:964;Ireland:353;Italy:39;Ivory Coast:225;Jamaica:1876;Japan:81;Jordan:962;Kazakhstan:7;Kenya:254;Kuwait:965;Kyrgyzstan:996;Laos:856;Latvia:371;Lebanon:961;Libya:218;Lithuania:370;Luxembourg:352;Macau:853;Madagascar:261;Malawi:265;Malaysia:60;Maldives:960;Mali:223;Malta:356;Mauritius:230;Mexico:52;Moldova:373;Mongolia:976;Montenegro:382;Morocco:212;Mozambique:258;Myanmar:95;Namibia:264;Nepal:977;Netherlands:31;New Zealand:64;Nicaragua:505;Niger:227;Nigeria:234;North Korea:850;North Macedonia:389;Norway:47;Oman:968;Pakistan:92;Palestine:970;Panama:507;Paraguay:595;Peru:51;Philippines:63;Poland:48;Portugal:351;Qatar:974;Romania:40;Russia:7;Rwanda:250;Saudi Arabia:966;Senegal:221;Serbia:381;Sierra Leone:232;Singapore:65;Slovakia:421;Slovenia:386;Somalia:252;South Africa:27;South Korea:82;South Sudan:211;Spain:34;Sri Lanka:94;Sudan:249;Sweden:46;Switzerland:41;Syria:963;Taiwan:886;Tajikistan:992;Tanzania:255;Thailand:66;Togo:228;Tunisia:216;Turkey:90;Turkmenistan:993;Uganda:256;Ukraine:380;United Arab Emirates:971;United Kingdom:44;United States:1;Uruguay:598;Uzbekistan:998;Venezuela:58;Vietnam:84;Yemen:967;Zambia:260;Zimbabwe:263".split(";")]
app.jinja_env.globals["country_lists"] = lambda: Markup(
    '<datalist id="countries">' + "".join(f'<option value="{n} (+{c})">' for n, c in COUNTRIES) + '</datalist>'
    '<datalist id="cnames">' + "".join(f'<option value="{n}">' for n, _ in COUNTRIES) + '</datalist>')

EMP_TABS = [
    ("personal", "Personal", [("staff_code", "Employee ID / Staff code", "readonly"), ("name", "Full name", "text"), ("father_name", "Father / Husband name (optional)", "text"),
        ("gender", "Gender (optional)", "select"), ("dob", "Date of birth", "date"), ("marital_status", "Marital status (optional)", "select"), ("nationality", "Nationality", "country"),
        ("blood_group", "Blood group (optional)", "select"), ("email", "Personal email", "email"), ("phone", "Mobile number (pick country, then number)", "phone"),
        ("whatsapp", "WhatsApp number (optional)", "phone")]),
    ("address", "Identity & Address", [("cnic", "CNIC (numbers only)", "cnic"), ("passport", "Passport / ID no. (international)", "text"), ("address", "Current address", "textarea"),
        ("permanent_address", "Permanent address", "textarea"), ("city", "City / District", "text"), ("province", "Province / State", "list:provinces"),
        ("postal_code", "Postal code (optional)", "text"), ("country", "Country", "country")]),
    ("quals", "Qualifications & Documents", [("skills", "Skills & qualifications", "textarea"), ("prev_employment", "Previous employment details", "textarea"),
        ("references_txt", "References (optional)", "textarea")]),
    ("job", "Job Details", [("join_date", "Joining date", "date"), ("branch", "Branch / Location", "text"), ("department", "Department", "dept"), ("designation", "Designation / Job title", "desig"),
        ("manager", "Reporting manager", "list:mgrs"), ("emp_type", "Employment type", "select"), ("probation_start", "Probation start date", "date"), ("probation_end", "Probation end date", "date"),
        ("confirmation_date", "Confirmation date", "date"), ("shift", "Work shift", "list:shifts"), ("work_location", "Work location", "text"), ("status", "Employee status", "select")]),
    ("salary", "Salary & Bank", [("salary", "Basic salary", "number"), ("allowances", "Allowances", "number"), ("house_rent", "House rent allowance", "number"),
        ("medical_allowance", "Medical allowance", "number"), ("bonus", "Bonuses / incentives", "number"), ("ot_rate", "Overtime rate per hour (blank = standard formula)", "number"),
        ("pay_method", "Payment method", "select"), ("bank_name", "Bank name", "text"), ("account_title", "Account title", "text"), ("account_no", "Account number / IBAN", "text"),
        ("salary_freq", "Salary frequency", "select"), ("tax_deduction", "Tax / deductions (per month)", "number")]),
    ("emergency", "Emergency Contact", [("emergency_name", "Contact person name", "text"), ("emergency_rel", "Relationship", "text"), ("emergency_phone", "Contact number", "phone"),
        ("emergency_alt", "Alternate number (optional)", "phone"), ("emergency_address", "Address (optional)", "textarea")]),
]
EMP_FIELDS = [f for _, _, fl in EMP_TABS for f in fl if f[2] != "readonly"]
# kind: (label, accepted types, multiple files, images only)
_IMG, _DOC = "image/*", ".pdf,.jpg,.jpeg,.png,.webp"
FILE_KINDS = {"photo": ("Profile picture", _IMG, False, True), "cnic_front": ("CNIC front image", _IMG, False, True), "cnic_back": ("CNIC back image", _IMG, False, True),
    "cv": ("CV / Resume (PDF, DOCX)", ".pdf,.doc,.docx", True, False), "matric": ("Matric / Intermediate certificates", _DOC, True, False), "degree": ("Degree / Diploma", _DOC, True, False),
    "transcript": ("Transcript / Marksheets", _DOC, True, False), "certificate": ("Professional certificates", _DOC, True, False), "experience": ("Experience letters", _DOC, True, False),
    "contract": ("Contract / Appointment letter", ".pdf,.doc,.docx,.jpg,.jpeg,.png", True, False)}
TAB_FILES = {"personal": ["photo"], "address": ["cnic_front", "cnic_back"], "quals": ["cv", "matric", "degree", "transcript", "certificate", "experience"], "job": ["contract"]}
IMG_EXT = {".png", ".jpg", ".jpeg", ".webp", ".gif"}; DOC_EXT = IMG_EXT | {".pdf", ".doc", ".docx"}
PROVINCES = ["Punjab", "Sindh", "Khyber Pakhtunkhwa", "Balochistan", "Gilgit-Baltistan", "Azad Kashmir", "Islamabad Capital Territory"]
SHIFTS = ["General (09:00 - 18:00)", "Morning (06:00 - 14:00)", "Evening (14:00 - 22:00)", "Night (22:00 - 06:00)"]
SELECTS = {"role": ["admin", "hr", "employee"], "gender": ["", "Male", "Female", "Other"], "marital_status": ["", "Single", "Married", "Divorced", "Widowed"],
           "blood_group": ["", "A+", "A-", "B+", "B-", "AB+", "AB-", "O+", "O-"], "emp_type": ["Permanent", "Contract", "Intern", "Part-time"],
           "status": ["Active", "On Leave", "Resigned", "Terminated"], "pay_method": ["Bank Transfer", "Cash", "Cheque", "Online Wallet"],
           "salary_freq": ["Monthly", "Weekly", "Daily", "Fortnightly"]}
app.jinja_env.globals.update(selects=SELECTS, tabs=EMP_TABS, tabfiles=TAB_FILES, fkinds=FILE_KINDS)
app.jinja_env.globals["emp_lists"] = lambda: Markup('<datalist id="provinces">' + "".join(f'<option value="{x}">' for x in PROVINCES) + '</datalist><datalist id="shifts">' + "".join(f'<option value="{x}">' for x in SHIFTS) + '</datalist>')

BULK = """<div class="bulk" id="bulk"><span><b id="bn">0</b> selected</span>
<select id="bkind"><option value="att">Attendance report</option><option value="info">Employee information</option></select>
<span class="o" id="bopt"><input type="month" id="bm" value="{{bmonth}}"><label><input type="checkbox" id="bd" checked> Daily details</label></span>
<select id="bf"><option value="pdf">PDF</option><option value="xlsx">Excel</option><option value="docx">Word</option></select>
<button type="button" id="bdl">{{icon('download')}} Download</button><button type="button" id="bdel" class="danger">Delete selected</button><button type="button" id="bclr" class="ghost">Clear</button></div>"""

FM = """{% macro field(k, label, t, val='') %}<div class="fld{{' wide' if t=='textarea'}}"><label>{{label}}</label>
{% if t=='readonly' %}<input value="{{val or 'Automatic'}}" disabled>
{% elif t=='textarea' %}<textarea name="{{k}}" rows="2">{{val or ''}}</textarea>
{% elif t=='phone' %}<div class="ph"><input class="cc" list="countries" placeholder="Search country" autocomplete="off"><input class="pn" type="tel" inputmode="numeric" placeholder="Phone number"><input type="hidden" name="{{k}}" value="{{val or ''}}"></div>
{% elif t=='cnic' %}<input class="cnic" name="{{k}}" value="{{val or ''}}" placeholder="13 digits, dashes added automatically" inputmode="numeric" autocomplete="off">
{% elif t=='select' %}<select name="{{k}}">{% for o in selects[k] %}<option value="{{o}}" {{'selected' if val==o}}>{{o or '- Select -'}}</option>{% endfor %}</select>
{% elif t=='country' %}<input name="{{k}}" list="cnames" value="{{val or ''}}" placeholder="Search country" autocomplete="off">
{% elif t=='dept' %}<input name="{{k}}" list="depts" value="{{val or ''}}" placeholder="e.g. Stitching" autocomplete="off">
{% elif t=='desig' %}<input name="{{k}}" list="desigs" value="{{val or ''}}" placeholder="e.g. Operator" autocomplete="off">
{% elif t.startswith('list:') %}<input name="{{k}}" list="{{t[5:]}}" value="{{val or ''}}" autocomplete="off">
{% else %}<input name="{{k}}" type="{{t}}" step="any" value="{{(val|int if (val is number and val == val|int) else val) if val is not none else ''}}" {{'required' if k=='name'}}>{% endif %}</div>{% endmacro %}
{% macro tabform(get) %}<div class="tabs">{% for tid,title,flds in tabs %}<button type="button" class="tab{{' on' if loop.first}}" data-t="{{tid}}">{{title}}</button>{% endfor %}</div>
{% for tid,title,flds in tabs %}<div class="panel{{' on' if loop.first}}" id="p-{{tid}}"><div class="fgrid">{% for k,label,t in flds %}{{field(k,label,t,get(k))}}{% endfor %}
{% for kd in tabfiles.get(tid, []) %}<div class="fld"><label>{{fkinds[kd][0]}}</label><input type="file" name="f_{{kd}}" accept="{{fkinds[kd][1]}}" {{'multiple' if fkinds[kd][2]}}></div>{% endfor %}</div></div>{% endfor %}{% endmacro %}
{{country_lists()}}{{emp_lists()}}<datalist id="depts">{% for x in depts %}<option value="{{x}}">{% endfor %}</datalist><datalist id="desigs">{% for x in desigs %}<option value="{{x}}">{% endfor %}</datalist><datalist id="mgrs">{% for x in names %}<option value="{{x}}">{% endfor %}</datalist>"""

def fmt_cnic(v):
    dg = re.sub(r"\D", "", v or "")
    return f"{dg[:5]}-{dg[5:12]}-{dg[12]}" if len(dg) == 13 else None

def lists(d):
    q = lambda col: [r[0] for r in d.execute(f"SELECT DISTINCT {col} FROM employees WHERE active=1 AND {col}<>'' ORDER BY {col}").fetchall()]
    return q("department"), q("designation")

def emp_values(f, d, exclude=None):
    """Read and validate the employee form. Returns (values, error)."""
    raw = f.get("cnic", "").strip(); cn = fmt_cnic(raw)
    if raw and not cn: return None, "CNIC must contain exactly 13 digits."
    if cn and d.execute("SELECT 1 FROM employees WHERE cnic=? AND active=1" + (" AND id<>?" if exclude else ""), (cn,) + ((exclude,) if exclude else ())).fetchone():
        return None, "An employee with this CNIC already exists."
    vals = {}
    try:
        for k, _, t in EMP_FIELDS:
            v = f.get(k, "").strip()
            vals[k] = (cn or "") if k == "cnic" else ((float(v) if v else None) if t == "number" else v)
    except ValueError: return None, "Please enter valid numbers in the salary fields."
    if not vals["name"]: return None, "Full name is required."
    vals["salary"] = vals["salary"] or 0; vals["status"] = vals["status"] or "Active"; vals["emp_type"] = vals["emp_type"] or "Permanent"
    return vals, None

def save_uploads(d, eid, files):
    saved, bad = 0, []
    for kind, (label, _, multi, img) in FILE_KINDS.items():
        for f in files.getlist("f_" + kind):
            if not f or not f.filename: continue
            ext = os.path.splitext(f.filename)[1].lower()
            if ext not in (IMG_EXT if img else DOC_EXT): bad.append(f.filename); continue
            if not multi:   # single-file kinds replace the old file
                for r in d.execute("SELECT id, stored FROM documents WHERE emp_id=? AND kind=?", (eid, kind)).fetchall():
                    try: os.remove(os.path.join(UPLOADS, r["stored"]))
                    except OSError: pass
                    d.execute("DELETE FROM documents WHERE id=?", (r["id"],))
            stored = uuid.uuid4().hex + ext; f.save(os.path.join(UPLOADS, stored))
            d.execute("INSERT INTO documents(emp_id,kind,stored,orig,created) VALUES(?,?,?,?,?)", (eid, kind, stored, secure_filename(f.filename)[:120] or "file" + ext, date.today().isoformat()))
            saved += 1
    return saved, bad

def photo_path(eid):
    r = db().execute("SELECT stored FROM documents WHERE emp_id=? AND kind='photo' ORDER BY id DESC", (eid,)).fetchone()
    return os.path.join(UPLOADS, r["stored"]) if r else None

def employees_view(d, form=None):
    rows = d.execute("SELECT * FROM employees WHERE active=1 ORDER BY id DESC").fetchall(); depts, desigs = lists(d)
    photos = {r[0] for r in d.execute("SELECT DISTINCT emp_id FROM documents WHERE kind='photo'").fetchall()}
    today = date.today().isoformat(); get = (lambda k: form.get(k, "")) if form else (lambda k: today if k == "join_date" else "")
    return page(FM + """<div class="phead"><h2>Employees</h2><button type="button" onclick="var b=document.getElementById('addbox');b.style.display=b.style.display=='none'?'block':'none'">{{icon('plus')}} Add employee</button></div>
    <div class="card" id="addbox" style="display:{{'block' if form else 'none'}}"><h3>New employee</h3><form method="post" enctype="multipart/form-data">{{tabform(get)}}
    <p style="margin-top:22px"><button>Save employee</button></p></form></div>
    <div class="card"><table class="ft bt" data-ph="Search by name, father / husband name, CNIC, phone, staff code or ID...">
    <tr><th><input type="checkbox" class="all" title="Select all"></th><th>ID</th><th>Code</th><th>Employee</th><th>CNIC</th><th>Phone</th><th>Department</th><th>Designation</th><th>Status</th><th>Basic salary</th><th>Joined</th><th></th></tr>
    {% for r in rows %}<tr><td><input type="checkbox" class="rc" value="{{r.id}}"></td><td class="idc">{{r.id}}</td><td class="nw">{{r.staff_code or ''}}</td>
    <td><div class="nm"><span class="av2">{% if r.id in photos %}<img src="/emp/{{r.id}}/photo" alt="">{% else %}{{(r.name or '?')[:1]|upper}}{% endif %}</span><div><a href="/employee/{{r.id}}">{{r.name}}</a>{% if r.father_name %}<small>Father / Husband: {{r.father_name}}</small>{% endif %}</div></div></td>
    <td class="num">{{r.cnic or ''}}{% if r.passport %}<small>Passport: {{r.passport}}</small>{% endif %}</td><td class="num">{{r.phone or ''}}</td>
    <td>{{r.department}}</td><td>{{r.designation}}</td><td><span class="bd b-{{(r.status or 'Active')|replace(' ','')}}">{{r.status or 'Active'}}</span></td>
    <td class="nw">{{'{:,.0f}'.format(r.salary or 0)}}</td><td class="nw">{{r.join_date}}</td>
    """ + act("employees") + """</tr>{% endfor %}</table></div>""" + BULK,
        rows=rows, form=form, get=get, depts=depts, desigs=desigs, names=[r["name"] for r in rows], photos=photos, bmonth=today[:7])

@app.route("/employees", methods=["GET", "POST"])
@login_required("admin", "hr")
def employees():
    d = db()
    if request.method == "POST":
        vals, err = emp_values(request.form, d)
        if err:
            flash(err); return employees_view(d, request.form)
        cols = list(vals)
        eid = insert("INSERT INTO employees(" + ",".join(cols) + ") VALUES(" + ",".join("?" * len(cols)) + ")", [vals[k] for k in cols])
        d.execute("UPDATE employees SET staff_code=? WHERE id=?", (f"EMP-{eid:04d}", eid))
        d.execute("INSERT INTO salary_history(emp_id,day,basic,note) VALUES(?,?,?,?)", (eid, date.today().isoformat(), vals["salary"], "Initial salary"))
        n, bad = save_uploads(d, eid, request.files); d.commit()
        flash(f"Employee added (staff code EMP-{eid:04d})." + (f" {n} file(s) uploaded." if n else "") + (" Skipped unsupported files: " + ", ".join(bad) if bad else "")); return redirect("/employees")
    return employees_view(d)

def employee_form_page(d, e, form=None):
    depts, desigs = lists(d); get = (lambda k: form.get(k, "")) if form else (lambda k: e[k])
    docs = d.execute("SELECT * FROM documents WHERE emp_id=? ORDER BY id", (e["id"],)).fetchall()
    return page(FM + """<div class="phead"><h2>Edit employee</h2><a href="/employee/{{e.id}}">View profile</a></div>
    <div class="card"><form method="post" enctype="multipart/form-data"><div class="fgrid" style="margin-bottom:20px"><div class="fld"><label>Employee ID (number)</label><input name="new_id" type="number" min="1" value="{{e['id']}}" required></div></div>
    {{tabform(get)}}<p style="margin-top:22px"><button>Save changes</button> <a class="cx" href="/employees">Cancel</a></p></form></div>
    {% if docs %}<div class="card"><h3>Files on record</h3><p style="color:var(--mut)">Uploading a new Profile picture or CNIC image replaces the old one. Delete other files here.</p>
    {% for x in docs %}<div class="doc"><b>{{fkinds[x.kind][0]}}</b><a href="/file/{{x.id}}" target="_blank">{{x.orig}}</a><form method="post" action="/doc/delete/{{x.id}}" onsubmit="return confirm('Delete this file?')"><button>Delete</button></form></div>{% endfor %}</div>{% endif %}""",
        e=e, get=get, depts=depts, desigs=desigs, names=[r["name"] for r in d.execute("SELECT name FROM employees WHERE active=1 ORDER BY name").fetchall()], docs=docs)

@app.route("/edit/employees/<int:i>", methods=["GET", "POST"])
@login_required("admin", "hr")
def employee_edit(i):
    d = db(); e = d.execute("SELECT * FROM employees WHERE id=? AND active=1", (i,)).fetchone()
    if not e: return "Not found", 404
    if request.method == "POST":
        f = request.form; vals, err = emp_values(f, d, exclude=i)
        try:
            new_id = int(f.get("new_id", "")); assert new_id >= 1
        except (ValueError, AssertionError): err = err or "Employee ID must be 1 or greater."; new_id = i
        if not err and new_id != i and d.execute("SELECT 1 FROM employees WHERE id=?", (new_id,)).fetchone(): err = f"ID {new_id} already belongs to another employee."
        if err:
            flash(err); return employee_form_page(d, e, f)
        d.execute("UPDATE employees SET " + ",".join(k + "=?" for k in vals) + " WHERE id=?", list(vals.values()) + [i])
        old = float(e["salary"] or 0)
        if vals["salary"] != old:
            d.execute("INSERT INTO salary_history(emp_id,day,basic,note) VALUES(?,?,?,?)", (i, date.today().isoformat(), vals["salary"], f"Changed from {old:,.0f} to {vals['salary']:,.0f}"))
        n, bad = save_uploads(d, i, request.files)
        if new_id != i:
            for t in ("attendance", "leaves", "overtime", "loans", "documents", "salary_history", "corrections", "users"): d.execute(f"UPDATE {t} SET emp_id=? WHERE emp_id=?", (new_id, i))
            d.execute("UPDATE employees SET id=? WHERE id=?", (new_id, i))
            if DATABASE_URL: d.execute("SELECT setval(pg_get_serial_sequence('employees','id'), COALESCE((SELECT MAX(id) FROM employees),0)+1, false)")
        d.commit(); flash("Changes saved." + (f" {n} file(s) uploaded." if n else "") + (" Skipped unsupported files: " + ", ".join(bad) if bad else "")); return redirect(f"/employee/{new_id}")
    return employee_form_page(d, e)

# ---------- Documents / photos ----------
def can_see(eid): return session.get("role") in ("admin", "hr") or session.get("emp_id") == eid

@app.route("/file/<int:i>")
@login_required()
def file_get(i):
    r = db().execute("SELECT * FROM documents WHERE id=?", (i,)).fetchone()
    if not r: return "Not found", 404
    if not can_see(r["emp_id"]): return "Access denied", 403
    return send_file(os.path.join(UPLOADS, r["stored"]), download_name=r["orig"], as_attachment=request.args.get("dl") == "1")

@app.route("/emp/<int:eid>/photo")
@login_required()
def emp_photo(eid):
    if not can_see(eid): return "Access denied", 403
    p = photo_path(eid)
    return send_file(p) if p and os.path.exists(p) else ("", 404)

@app.route("/doc/delete/<int:i>", methods=["POST"])
@login_required("admin", "hr")
def doc_delete(i):
    d = db(); r = d.execute("SELECT * FROM documents WHERE id=?", (i,)).fetchone()
    if r:
        try: os.remove(os.path.join(UPLOADS, r["stored"]))
        except OSError: pass
        d.execute("DELETE FROM documents WHERE id=?", (i,)); d.commit(); flash("File deleted.")
    return redirect(f"/employee/{r['emp_id']}" if r else "/employees")

@app.route("/employee/<int:i>")
@login_required("admin", "hr")
def employee_profile(i):
    d = db(); e = d.execute("SELECT * FROM employees WHERE id=? AND active=1", (i,)).fetchone()
    if not e: return "Not found", 404
    m = date.today().isoformat()[:7]; c = calc(d, e, m); clean = lambda l: re.sub(r"\s*\(.*?\)", "", l)
    def val(k, t):
        v = e[k]
        return "" if v in (None, "") else (f"{float(v):,.0f}" if t == "number" else v)
    secs = [(title, [(clean(l), val(k, t), "tel" if t == "phone" else "mail" if t == "email" else "wide" if t == "textarea" else "") for k, l, t in fl]) for _, title, fl in EMP_TABS]
    bydoc = {}
    for x in d.execute("SELECT * FROM documents WHERE emp_id=? ORDER BY id", (i,)).fetchall(): bydoc.setdefault(x["kind"], []).append(x)
    return page("""<div class="phead"><a href="/employees">&larr; Back to employees</a><span><a href="/bulk/download?ids={{e.id}}&amp;kind=info&amp;fmt=pdf">Download PDF</a> &nbsp;&middot;&nbsp; <a href="/edit/employees/{{e.id}}">Edit employee</a></span></div>
    <div class="card"><div class="pgh"><span class="av2">{% if bydoc.get('photo') %}<img src="/emp/{{e.id}}/photo" alt="">{% else %}{{(e.name or '?')[:1]|upper}}{% endif %}</span>
    <div><h2>{{e.name}}</h2><p>{{e.designation}}{% if e.department %} &middot; {{e.department}}{% endif %} &middot; ID {{e.id}} &middot; {{e.staff_code}} &nbsp;<span class="bd b-{{(e.status or 'Active')|replace(' ','')}}">{{e.status or 'Active'}}</span></p></div></div>
    <div class="pg"><div><small>{{mlabel}} attendance</small>{{c.present}} present, {{c.absent}} absent, {{c.half}} half day, {{c.late}} late, {{c.leave}} leave</div><div><small>Overtime this month</small>{{c.ot_hours}} hours</div><div><small>Net pay this month</small>{{'{:,.0f}'.format(c.net)}}</div>
    {% for k,v in bal.items() %}<div><small>{{k}} leave balance</small>{{v.left}} of {{v.quota}} days left</div>{% endfor %}</div></div>
    {% for title,items in secs %}<div class="card"><h3>{{title}}</h3><div class="pg">{% for l,v,k in items %}<div{{' style="grid-column:1/-1"'|safe if k=='wide'}}><small>{{l}}</small>{% if v %}{% if k=='tel' %}<a href="tel:{{v|replace(' ','')}}">{{v}}</a>{% elif k=='mail' %}<a href="mailto:{{v}}">{{v}}</a>{% else %}{{v}}{% endif %}{% else %}<span style="color:#aab3c2">-</span>{% endif %}</div>{% endfor %}</div>
    {% if title=='Salary & Bank' %}<h3 style="margin-top:22px">Salary history</h3><table><tr><th>Date</th><th>Basic salary</th><th>Note</th></tr>{% for h in hist %}<tr><td>{{h.day}}</td><td>{{'{:,.0f}'.format(h.basic or 0)}}</td><td>{{h.note}}</td></tr>{% endfor %}</table>{% endif %}</div>{% endfor %}
    <div class="card"><h3>Documents</h3>{% if not bydoc %}<p style="color:var(--mut)">No documents uploaded yet. Use Edit employee to upload files.</p>{% endif %}
    <div class="pg">{% for kd,fk in fkinds.items() %}{% if bydoc.get(kd) %}<div><small>{{fk[0]}}</small>{% for x in bydoc[kd] %}<div class="doc">{% if fk[3] %}<a href="/file/{{x.id}}" target="_blank"><img class="thumb" src="/file/{{x.id}}" alt=""></a>{% else %}<a href="/file/{{x.id}}" target="_blank">{{x.orig}}</a>{% endif %}
    <a href="/file/{{x.id}}?dl=1">Download</a><form method="post" action="/doc/delete/{{x.id}}" onsubmit="return confirm('Delete this file?')"><button>Delete</button></form></div>{% endfor %}</div>{% endif %}{% endfor %}</div></div>""",
        e=e, c=c, secs=secs, bydoc=bydoc, mlabel=month_label(m), bal=leave_balance(d, i), hist=d.execute("SELECT * FROM salary_history WHERE emp_id=? ORDER BY id DESC", (i,)).fetchall())

# ---------- Attendance (monthly chart, click a cell to mark) ----------
STATUSES = ("Present", "Absent", "Late", "Leave", "Half Day")
SHIFT_END = "18:00"      # check-out before this counts as early departure
WEEKLY_OFF = (6,)        # 6 = Sunday (0 = Monday). Example for Friday + Sunday: (4, 6)
WD = ["Mo", "Tu", "We", "Th", "Fr", "Sa", "Su"]
SHIFT_START = "09:00"   # shift start time; check-in after this (plus grace) counts as late
GRACE_MIN = 0           # minutes of grace before a check-in is called Late

def tmin(t):
    try: h, m = (t or "").split(":"); return int(h) * 60 + int(m)
    except Exception: return None

def clean_time(t):
    t = (t or "").strip(); v = tmin(t)
    return t[:5] if v is not None and 0 <= v < 1440 else None

def worked(cin, cout):
    a, b = tmin(cin), tmin(cout)
    return 0.0 if a is None or b is None else round(((b - a) % 1440) / 60, 2)

@app.route("/attendance")
@login_required("admin", "hr")
def attendance():
    d = db(); today = date.today(); m = request.args.get("month", "")
    try: y, mo = int(m[:4]), int(m[5:7]); first = date(y, mo, 1)
    except Exception: y, mo = today.year, today.month; first = date(y, mo, 1)
    m = first.isoformat()[:7]
    holmap = {r["day"]: r["name"] for r in d.execute("SELECT day, name FROM holidays WHERE day LIKE ?", (m + "%",)).fetchall()}
    days = [dict(n=i, iso=f"{m}-{i:02d}", wd=WD[date(y, mo, i).weekday()], sun=date(y, mo, i).weekday() in WEEKLY_OFF, today=date(y, mo, i) == today, hol=holmap.get(f"{m}-{i:02d}"))
            for i in range(1, calendar.monthrange(y, mo)[1] + 1)]
    emps = d.execute("SELECT id, name, department FROM employees WHERE active=1 AND COALESCE(status,'Active') NOT IN ('Resigned','Terminated') ORDER BY id").fetchall()
    got = {(r["emp_id"], r["day"]): r for r in d.execute("SELECT emp_id, day, status, check_in, check_out, late_min, early_min FROM attendance WHERE day LIKE ?", (m + "%",)).fetchall()}
    rows = []
    for e in emps:
        cells = []
        for x in days:
            r = got.get((e["id"], x["iso"])); st = r["status"] if r else ""; cin = (r["check_in"] or "") if r else ""; cout = (r["check_out"] or "") if r else ""
            late = r["late_min"] if r and r["late_min"] is not None else ""; early = r["early_min"] if r and r["early_min"] is not None else ""
            tip = " | ".join(p for p in [("In " + cin) if cin else "", ("Out " + cout) if cout else "", f"Late {late} min" if late else "", f"Left early {early} min" if early else "",
                                          f"{worked(cin, cout):g} h" if cin and cout else ""] if p)
            cells.append(dict(k=SYMBOL.get(st, ""), cin=cin, cout=cout, late=late, early=early, tip=tip))
        tot = [sum(1 for c in cells if c["k"] == k) for k in "PALVH"]
        rows.append(dict(e=e, cells=cells, tot=tot))
    tc = {k: d.execute("SELECT COUNT(*) FROM attendance a JOIN employees e ON e.id=a.emp_id WHERE e.active=1 AND COALESCE(e.status,'Active') NOT IN ('Resigned','Terminated') AND a.day=? AND a.status=?", (today.isoformat(), k)).fetchone()[0] for k in STATUSES}
    tc["N"] = max(len(emps) - sum(tc.values()), 0)
    return page("""<style>
    .mnav{display:flex;gap:8px;align-items:center;flex-wrap:wrap}.mnav a{display:inline-flex;align-items:center;justify-content:center;min-width:38px;height:40px;padding:0 12px;border:1px solid var(--line);border-radius:10px;background:#fff;color:var(--ink);font-weight:700;text-decoration:none}
    .mnav a:hover{background:var(--navy);color:#fff}.mnav input{height:40px}
    .kpis{display:grid;grid-template-columns:repeat(5,1fr);gap:14px;margin-bottom:20px}@media(max-width:900px){.kpis{grid-template-columns:repeat(2,1fr)}}
    .kp{background:#fff;border:1px solid var(--line);border-radius:14px;padding:14px 18px;box-shadow:var(--shadow);border-left:5px solid var(--c)}.kp b{display:block;font-size:26px;font-weight:800;letter-spacing:-.5px}.kp span{color:var(--mut);font-size:12.5px;font-weight:600}
    .legend{display:flex;gap:14px;flex-wrap:wrap;margin-left:auto;font-size:12.5px;font-weight:600;color:var(--mut)}.legend i,.dot{display:inline-block;width:12px;height:12px;border-radius:4px;margin-right:6px;vertical-align:-1px}
    .st-P{background:#e3f6ec;color:#12875b}.st-A{background:#fde8e8;color:#c0362c}.st-L{background:#fff1d6;color:#b9770e}.st-V{background:#e8efff;color:#2f5fe3}.st-H{background:#efe8ff;color:#6d3fd1}
    .attwrap{overflow:auto;max-height:68vh;border:1px solid var(--line);border-radius:14px;background:#fff}
    table.att{border-collapse:separate;border-spacing:0;width:max-content;min-width:100%}
    table.att th,table.att td{padding:0;border-bottom:1px solid #eef1f6;border-right:1px solid #f3f5f9;text-align:center;background:none}
    table.att thead th{position:sticky;top:0;z-index:3;background:#f7f9fc;height:50px;min-width:38px;font-size:13px;color:var(--ink);text-transform:none;letter-spacing:0;border-radius:0}
    table.att thead th small{display:block;font-size:10.5px;color:var(--mut);font-weight:600}
    table.att thead th.dh{cursor:pointer}table.att thead th.dh:hover{background:#eaeff8}
    table.att thead th.sun,td.cell.sun{background:#f1f3f8}table.att thead th.hol,td.cell.hol{background:#fbe9d6}table.att thead th.today{background:#ffe8b8;color:#8a5600}
    table.att .nm{position:sticky;left:0;z-index:2;background:#fff;text-align:left;min-width:270px;padding:10px 16px;border-right:1px solid var(--line)}table.att thead th.nm{z-index:4;background:#f7f9fc}
    table.att tr:hover td.nm{background:#fafbfe}
    td.cell{cursor:pointer;height:48px;min-width:38px;font-weight:800;font-size:12.5px;transition:.1s}td.cell:hover{outline:2px solid var(--acc);outline-offset:-2px}
    td.cell sub{display:block;font-size:9.5px;line-height:1;margin-top:1px;font-weight:700}
    td.cell:empty::after{content:"\\00b7";color:#cbd3e1;font-size:18px}td.cell.tcol{box-shadow:inset 2px 0 0 #f5a524,inset -2px 0 0 #f5a524}
    table.att td.st-P,table.att th.st-P{background:#e3f6ec;color:#12875b}table.att td.st-A,table.att th.st-A{background:#fde8e8;color:#c0362c}
    table.att td.st-L,table.att th.st-L{background:#fff1d6;color:#b9770e}table.att td.st-V,table.att th.st-V{background:#e8efff;color:#2f5fe3}table.att td.st-H,table.att th.st-H{background:#efe8ff;color:#6d3fd1}
    table.att th.t,table.att td.t{position:sticky;min-width:44px;width:44px;z-index:2}table.att thead th.t{z-index:4}td.t{font-weight:800;background:#fff}
    table.att .t:nth-last-child(1){right:0}table.att .t:nth-last-child(2){right:44px}table.att .t:nth-last-child(3){right:88px}table.att .t:nth-last-child(4){right:132px}table.att .t:nth-last-child(5){right:176px;box-shadow:-6px 0 8px -6px rgba(15,27,45,.18)}
    .pop{position:fixed;display:none;z-index:50;background:#fff;border:1px solid var(--line);border-radius:14px;box-shadow:0 18px 50px rgba(15,27,45,.22);padding:8px;min-width:250px}
    .pop .pt{font-size:11px;text-transform:uppercase;letter-spacing:.8px;color:var(--mut);font-weight:700;padding:6px 10px}
    .pop .tm{display:grid;grid-template-columns:1fr 1fr;gap:8px;padding:0 10px 8px}.pop .tm label{font-size:11.5px;font-weight:700;color:var(--mut)}.pop .tm input{width:100%;padding:7px 8px;margin-top:3px;font-size:13px}
    .pop button{display:flex;align-items:center;width:100%;gap:8px;background:none;color:var(--ink);box-shadow:none;border-radius:9px;padding:9px 10px;text-align:left;font-weight:600}.pop button:hover{background:#f3f6fc;transform:none}
    .pop kbd{margin-left:auto;font:600 11px monospace;color:var(--mut);background:#f1f3f8;border-radius:5px;padding:2px 6px}.pop .clr{color:var(--red);border-top:1px solid var(--line);border-radius:0 0 9px 9px;margin-top:4px}
    .toast{position:fixed;right:24px;bottom:24px;background:var(--navy);color:#fff;padding:11px 18px;border-radius:12px;font-weight:600;opacity:0;transform:translateY(10px);transition:.2s;pointer-events:none;z-index:60}.toast.show{opacity:1;transform:none}.toast.bad{background:var(--red)}
    .hint{color:var(--mut);font-size:12.5px;margin:12px 2px 0}
    </style>
    <div class="phead"><h2>Attendance</h2><div class="mnav"><a href="/attendance?month={{prev}}" title="Previous month">&lsaquo;</a><input type="month" id="mo" value="{{m}}"><a href="/attendance?month={{nxt}}" title="Next month">&rsaquo;</a>
    <a href="/attendance">Today</a><a href="/attendance/report.xlsx?month={{m}}">{{icon('download')}}&nbsp;Excel</a></div></div>
    <div class="kpis"><div class="kp" style="--c:#12875b"><b id="k-Present">{{tc.Present}}</b><span>Present today</span></div><div class="kp" style="--c:#c0362c"><b id="k-Absent">{{tc.Absent}}</b><span>Absent today</span></div>
    <div class="kp" style="--c:#e0900f"><b id="k-Late">{{tc.Late}}</b><span>Late today</span></div><div class="kp" style="--c:#2f5fe3"><b id="k-Leave">{{tc.Leave}}</b><span>On leave today</span></div>
    <div class="kp" style="--c:#8794ab"><b id="k-N">{{tc.N}}</b><span>Not marked today</span></div></div>
    {% if not rows %}<div class="card">No employees yet. Add employees first, then mark their attendance here.</div>{% else %}
    <div class="card"><div class="tb"><div class="s">{{icon('search')}}<input id="q" type="search" placeholder="Search employee, ID or department..."></div>
    <select id="ds"><option value="">All departments</option>{% for x in depts %}<option>{{x}}</option>{% endfor %}</select><span class="cnt" id="cn"></span>
    <div class="legend"><span><i class="st-P"></i>Present</span><span><i class="st-A"></i>Absent</span><span><i class="st-L"></i>Late</span><span><i class="st-V"></i>Leave</span><span><i class="st-H"></i>Half day</span></div></div>
    <div class="attwrap"><table class="att bt"><thead><tr><th class="nm"><label style="display:flex;gap:10px;align-items:center"><input type="checkbox" class="all" title="Select all"> Employee</label></th>
    {% for x in days %}<th class="dh {{'sun' if x.sun}} {{'hol' if x.hol}} {{'today' if x.today}}" data-d="{{x.iso}}" title="{{x.hol or ('Weekly off' if x.sun else 'Click to mark everyone')}} - {{x.iso}}">{{x.n}}<small>{{x.wd}}</small></th>{% endfor %}
    <th class="t st-P">P</th><th class="t st-A">A</th><th class="t st-L">L</th><th class="t st-V">V</th><th class="t st-H">H</th></tr></thead><tbody>
    {% for r in rows %}<tr data-dept="{{r.e.department or ''}}" data-s="{{(r.e.name ~ ' ' ~ r.e.id ~ ' ' ~ (r.e.department or ''))|lower}}">
    <td class="nm"><div class="nmw" style="display:flex;align-items:center"><input type="checkbox" class="rc" value="{{r.e.id}}" style="margin-right:12px"><span class="av2">{{(r.e.name or '?')[:1]|upper}}</span><div><b>{{r.e.name}}</b><small>ID {{r.e.id}}{% if r.e.department %} &middot; {{r.e.department}}{% endif %}</small></div></div></td>
    {% for x in days %}{% set c = r.cells[loop.index0] %}<td class="cell {{'st-'~c.k if c.k}} {{'sun' if x.sun}} {{'hol' if x.hol}} {{'tcol' if x.today}}" data-e="{{r.e.id}}" data-d="{{x.iso}}" data-st="{{c.k}}" data-in="{{c.cin}}" data-out="{{c.cout}}" data-late="{{c.late}}" data-early="{{c.early}}" title="{{c.tip or x.hol or ''}}">{{c.k}}{% if c.k=='L' and c.late %}<sub>{{c.late}}</sub>{% endif %}</td>{% endfor %}
    {% for v in r.tot %}<td class="t">{{v}}</td>{% endfor %}</tr>{% endfor %}</tbody></table></div>
    <p class="hint">Click a cell to mark Present, Absent, Late, Leave or Half day. Optionally add check-in / check-out time first (check-in after {{shift}} is marked Late, check-out before {{shift_end}} is recorded as leaving early). Shortcuts: P / A / L / V / H. Click a date at the top to mark everyone not yet marked.</p></div>{% endif %}
    <div class="pop" id="pop"></div><div class="toast" id="toast"></div>
    <script>
    (function(){
    var tbl=document.querySelector('table.att'),SHIFT='{{shift}}',SHIFT_END='{{shift_end}}';
    document.getElementById('mo').addEventListener('change',function(){if(this.value)location='/attendance?month='+this.value;});
    if(!tbl)return;
    var S={P:'Present',A:'Absent',L:'Late',V:'Leave',H:'Half Day'},pop=document.getElementById('pop'),tg=document.getElementById('toast'),tgt=null,mode='';
    var rows=[].slice.call(tbl.tBodies[0].rows),q=document.getElementById('q'),ds=document.getElementById('ds'),cn=document.getElementById('cn');
    function toast(t,bad){tg.textContent=t;tg.className='toast show'+(bad?' bad':'');clearTimeout(toast.h);toast.h=setTimeout(function(){tg.className='toast';},1700);}
    function post(u,b){return fetch(u,{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify(b)}).then(function(r){return r.json();});}
    function mins(t){var p=(t||'').split(':');return p.length===2?(+p[0])*60+(+p[1]):null;}
    function filt(){var s=q.value.trim().toLowerCase(),d=ds.value,n=0;rows.forEach(function(r){var ok=(!d||r.getAttribute('data-dept')===d)&&(!s||r.getAttribute('data-s').indexOf(s)>-1);r.style.display=ok?'':'none';if(ok)n++;});cn.textContent='Showing '+n+' of '+rows.length;}
    q.addEventListener('input',filt);ds.addEventListener('change',filt);filt();
    function paint(c,r){var k={Present:'P',Absent:'A',Late:'L',Leave:'V','Half Day':'H'}[r.status]||'';c.className=c.className.replace(/ ?st-[PALVH]/g,'')+(k?' st-'+k:'');
     c.setAttribute('data-st',k);c.setAttribute('data-in',r.check_in||'');c.setAttribute('data-out',r.check_out||'');c.setAttribute('data-late',r.late_min==null?'':r.late_min);c.setAttribute('data-early',r.early_min==null?'':r.early_min);
     c.innerHTML=k+(k==='L'&&r.late_min?'<sub>'+r.late_min+'</sub>':'');var t=[];if(r.check_in)t.push('In '+r.check_in);if(r.check_out)t.push('Out '+r.check_out);if(r.late_min)t.push('Late '+r.late_min+' min');if(r.early_min)t.push('Left early '+r.early_min+' min');if(r.hours)t.push(r.hours+' h');c.title=t.join(' | ');}
    function recalc(r){var cs=r.querySelectorAll('td.cell'),t=r.querySelectorAll('td.t');['P','A','L','V','H'].forEach(function(k,i){var n=0;[].forEach.call(cs,function(c){if(c.getAttribute('data-st')===k)n++;});t[i].textContent=n;});}
    function kpi(){var th=tbl.querySelector('th.today');if(!th)return;var ix=th.cellIndex,c={P:0,A:0,L:0,V:0,H:0};rows.forEach(function(r){var v=r.cells[ix].getAttribute('data-st');if(c[v]!==undefined)c[v]++;});
     document.getElementById('k-Present').textContent=c.P;document.getElementById('k-Absent').textContent=c.A;document.getElementById('k-Late').textContent=c.L;document.getElementById('k-Leave').textContent=c.V;document.getElementById('k-N').textContent=rows.length-c.P-c.A-c.L-c.V-c.H;}
    function open(el,m,title){tgt=el;mode=m;var h='<div class="pt">'+title+'</div>';
     if(m==='cell')h+='<div class="tm"><label>Check-in<input type="time" id="pin"></label><label>Check-out<input type="time" id="pout"></label><label>Late (min)<input type="number" min="0" id="plate" placeholder="auto"></label><label>Left early (min)<input type="number" min="0" id="pearly" placeholder="auto"></label></div><div class="pt">Mark as</div>';
     [['P','Present'],['A','Absent'],['L','Late'],['V','Leave'],['H','Half Day']].forEach(function(x){h+='<button data-s="'+x[0]+'"><i class="dot st-'+x[0]+'"></i>'+x[1]+'<kbd>'+x[0]+'</kbd></button>';});
     if(m==='cell')h+='<button data-s="" class="clr">Clear mark</button>';pop.innerHTML=h;pop.style.display='block';
     if(m==='cell'){var pin=pop.querySelector('#pin'),plate=pop.querySelector('#plate');pin.value=el.getAttribute('data-in')||'';pop.querySelector('#pout').value=el.getAttribute('data-out')||'';plate.value=el.getAttribute('data-late')||'';pop.querySelector('#pearly').value=el.getAttribute('data-early')||'';
      var pout=pop.querySelector('#pout'),pe=pop.querySelector('#pearly');pin.addEventListener('input',function(){var a=mins(pin.value),b=mins(SHIFT);plate.value=(a!==null&&b!==null&&a>b)?(a-b):'';});pout.addEventListener('input',function(){var a=mins(pout.value),b=mins(SHIFT_END);pe.value=(a!==null&&b!==null&&a<b)?(b-a):'';});}
     var r=el.getBoundingClientRect(),w=pop.offsetWidth,ht=pop.offsetHeight,x=Math.min(r.left,window.innerWidth-w-12),y=r.bottom+6;if(y+ht>window.innerHeight-10)y=r.top-ht-6;
     pop.style.left=Math.max(8,x)+'px';pop.style.top=Math.max(8,y)+'px';}
    function close(){pop.style.display='none';}
    function choose(s){var c=tgt,m=mode,p;
     if(m==='cell'){var pin=pop.querySelector('#pin').value,pout=pop.querySelector('#pout').value,pl=pop.querySelector('#plate').value,pe=pop.querySelector('#pearly').value;close();var row=c.parentNode;
      p={emp_id:+c.getAttribute('data-e'),day:c.getAttribute('data-d'),status:s?S[s]:''};if(s&&(pin||pout||pl!==''||pe!=='')){p.check_in=pin;p.check_out=pout;p.late_min=pl;p.early_min=pe;}
      post('/attendance/mark',p).then(function(r){if(r.ok){paint(c,r);recalc(row);kpi();toast(r.status?'Marked '+r.status:'Mark cleared');}else toast(r.error||'Could not save',1);}).catch(function(){toast('Network error',1);});}
     else{close();if(s)post('/attendance/mark_all',{day:c.getAttribute('data-d'),status:S[s]}).then(function(r){if(r.ok)location.reload();else toast(r.error||'Could not save',1);}).catch(function(){toast('Network error',1);});}}
    tbl.addEventListener('click',function(e){var c=e.target.closest('td.cell'),h=e.target.closest('th.dh');
     if(c){e.stopPropagation();open(c,'cell',c.parentNode.querySelector('b').textContent+' - '+c.getAttribute('data-d'));}else if(h){e.stopPropagation();open(h,'col','Mark all unmarked on '+h.getAttribute('data-d'));}});
    pop.addEventListener('click',function(e){var b=e.target.closest('button');if(b)choose(b.getAttribute('data-s'));e.stopPropagation();});
    document.addEventListener('click',close);window.addEventListener('scroll',close,true);window.addEventListener('resize',close);
    document.addEventListener('keydown',function(e){if(pop.style.display!=='block')return;var k=e.key.toUpperCase();if(k==='ESCAPE'){close();return;}
     if(e.target.tagName==='INPUT')return;if('PALVH'.indexOf(k)>-1&&k.length===1)choose(k);else if((k==='DELETE'||k==='BACKSPACE')&&mode==='cell')choose('');});
    })();
    </script>""" + BULK, rows=rows, days=days, m=m, tc=tc, shift=SHIFT_START, shift_end=SHIFT_END, depts=sorted({e["department"] for e in emps if e["department"]}), bmonth=m,
        prev=(first - timedelta(days=1)).isoformat()[:7], nxt=(first + timedelta(days=32)).replace(day=1).isoformat()[:7])

@app.route("/attendance/mark", methods=["POST"])
@login_required("admin", "hr")
def attendance_mark():
    j = request.get_json(silent=True) or {}; d = db()
    try: eid = int(j["emp_id"]); day = date.fromisoformat(j["day"]).isoformat(); st = j.get("status", "")
    except Exception: return {"ok": False, "error": "Invalid data"}, 400
    if st not in STATUSES + ("",) or not d.execute("SELECT 1 FROM employees WHERE id=? AND active=1", (eid,)).fetchone():
        return {"ok": False, "error": "Invalid employee or status"}, 400
    old = d.execute("SELECT * FROM attendance WHERE emp_id=? AND day=?", (eid, day)).fetchone(); WORKED = ("Present", "Late", "Half Day")
    timed = any(k in j for k in ("check_in", "check_out", "late_min", "early_min"))
    def num(v):
        try: return max(int(v), 0)
        except (TypeError, ValueError): return None
    if timed: cin, cout, late, early = clean_time(j.get("check_in")), clean_time(j.get("check_out")), num(j.get("late_min")), num(j.get("early_min"))
    elif old and st in WORKED and old["status"] in WORKED: cin, cout, late, early = old["check_in"], old["check_out"], old["late_min"], old["early_min"]
    else: cin = cout = late = early = None
    if st not in WORKED: cin = cout = late = early = None
    if st == "Present" and not timed: late = None
    if timed and st in WORKED:
        if late is None and cin: late = max(0, (tmin(cin) or 0) - tmin(SHIFT_START))
        if early is None and cout and st != "Half Day": early = max(0, tmin(SHIFT_END) - (tmin(cout) or 0))   # a half day is expected to end early
        if st == "Present" and late and late > GRACE_MIN: st = "Late"   # arrived after shift start
        if st == "Present": late = None
        late = late or None; early = early or None
    if st: d.execute("""INSERT INTO attendance(emp_id,day,status,check_in,check_out,late_min,early_min) VALUES(?,?,?,?,?,?,?) ON CONFLICT(emp_id,day) DO UPDATE SET
        status=excluded.status, check_in=excluded.check_in, check_out=excluded.check_out, late_min=excluded.late_min, early_min=excluded.early_min""", (eid, day, st, cin, cout, late, early))
    else: d.execute("DELETE FROM attendance WHERE emp_id=? AND day=?", (eid, day))
    d.commit()
    return {"ok": True, "status": st, "check_in": cin, "check_out": cout, "late_min": late, "early_min": early, "hours": worked(cin, cout) or None}

@app.route("/attendance/mark_all", methods=["POST"])
@login_required("admin", "hr")
def attendance_mark_all():
    j = request.get_json(silent=True) or {}; d = db()
    try: day = date.fromisoformat(j["day"]).isoformat(); st = j["status"]
    except Exception: return {"ok": False, "error": "Invalid data"}, 400
    if st not in STATUSES: return {"ok": False, "error": "Invalid status"}, 400
    ids = [r["id"] for r in d.execute("SELECT id FROM employees WHERE active=1 AND COALESCE(status,'Active') NOT IN ('Resigned','Terminated') AND id NOT IN (SELECT emp_id FROM attendance WHERE day=?)", (day,)).fetchall()]
    for i in ids: d.execute("INSERT INTO attendance(emp_id,day,status) VALUES(?,?,?)", (i, day, st))
    d.commit(); flash(f"{len(ids)} employee(s) marked {st} on {day}." if ids else "Everyone is already marked for that day.")
    return {"ok": True}

# ---------- Expenses + report ----------
@app.route("/expenses", methods=["GET", "POST"])
@login_required("admin")
def expenses():
    d = db()
    if request.method == "POST":
        f = request.form
        d.execute("INSERT INTO expenses(day,category,description,amount) VALUES(?,?,?,?)",
                  (f["day"], f["category"], f["description"], f["amount"]))
        d.commit(); flash("Expense added."); return redirect("/expenses")
    month = request.args.get("month", date.today().isoformat()[:7])
    rows = d.execute("SELECT * FROM expenses WHERE day LIKE ? ORDER BY day DESC", (month + "%",)).fetchall()
    cats = d.execute("SELECT category, SUM(amount) t FROM expenses WHERE day LIKE ? GROUP BY category", (month + "%",)).fetchall()
    total = sum(r["amount"] for r in rows)
    return page("""<h2>Company expenses</h2><div class="card"><form class="row" method="post">
    <input type="date" name="day" value="{{today}}"><select name="category"><option>Electricity</option><option>Raw material</option>
    <option>Salaries</option><option>Maintenance</option><option>Transport</option><option>Other</option></select>
    <input name="description" placeholder="Description"><input name="amount" type="number" step="0.01" placeholder="Amount" required>
    <button>Add expense</button></form>
    <form class="row" method="get"><input type="month" name="month" value="{{month}}"><button>Show report</button></form>
    <h3>Report {{month}} - Total: {{'{:,.0f}'.format(total)}}</h3>
    <p>{% for c in cats %}<b>{{c.category}}</b>: {{'{:,.0f}'.format(c.t)}} &nbsp; {% endfor %}</p>
    <table class="ft" data-ph="Search expenses..."><tr><th>Date</th><th>Category</th><th>Description</th><th>Amount</th><th></th></tr>
    {% for r in rows %}<tr><td>{{r.day}}</td><td>{{r.category}}</td><td>{{r.description}}</td><td>{{'{:,.0f}'.format(r.amount)}}</td><td class="noprint"><a href="/edit/expenses/{{r.id}}">Edit</a> <form method="post" action="/delete/expenses/{{r.id}}" style="display:inline" onsubmit="return confirm('Are you sure you want to delete this?')"><button>Delete</button></form></td></tr>{% endfor %}</table></div>""",
        rows=rows, cats=cats, total=total, month=month, today=date.today().isoformat())

# ---------- Inventory ----------
@app.route("/inventory", methods=["GET", "POST"])
@login_required("admin")
def inventory():
    d = db()
    if request.method == "POST":
        f = request.form
        if f.get("action") == "adjust":
            d.execute("UPDATE inventory SET qty=qty+? WHERE id=?", (float(f["delta"]), f["id"]))
        else:
            d.execute("INSERT INTO inventory(item,sku,qty,unit,min_qty,price) VALUES(?,?,?,?,?,?)",
                      (f["item"], f["sku"], f["qty"] or 0, f["unit"], f["min_qty"] or 0, f["price"] or 0))
        d.commit(); return redirect("/inventory")
    rows = d.execute("SELECT * FROM inventory ORDER BY item").fetchall()
    value = sum(r["qty"] * r["price"] for r in rows)
    return page("""<h2>Inventory</h2><div class="card"><form class="row" method="post">
    <input name="item" placeholder="Item" required><input name="sku" placeholder="SKU"><input name="qty" type="number" step="any" placeholder="Qty">
    <input name="unit" placeholder="Unit (kg, pcs)"><input name="min_qty" type="number" step="any" placeholder="Min qty"><input name="price" type="number" step="any" placeholder="Unit price">
    <button>Add item</button></form><p>Stock value: <b>{{'{:,.0f}'.format(value)}}</b></p>
    <table class="ft" data-ph="Search items or SKU..."><tr><th>Item</th><th>SKU</th><th>Qty</th><th>Min</th><th>Price</th><th>Stock in/out (+/-)</th><th></th></tr>
    {% for r in rows %}<tr><td>{{r.item}}</td><td>{{r.sku}}</td><td class="{{'low' if r.qty<=r.min_qty}}">{{r.qty}} {{r.unit}}</td><td>{{r.min_qty}}</td><td>{{r.price}}</td>
    <td><form method="post" class="row" style="margin:0"><input type="hidden" name="action" value="adjust"><input type="hidden" name="id" value="{{r.id}}">
    <input name="delta" type="number" step="any" style="width:80px" required><button>Update</button></form></td><td class="noprint"><a href="/edit/inventory/{{r.id}}">Edit</a> <form method="post" action="/delete/inventory/{{r.id}}" style="display:inline" onsubmit="return confirm('Are you sure you want to delete this?')"><button>Delete</button></form></td></tr>{% endfor %}</table></div>""",
        rows=rows, value=value)

# ---------- Users (admin only) ----------
@app.route("/users", methods=["GET", "POST"])
@login_required("admin")
def users():
    d = db()
    if request.method == "POST":
        try:
            d.execute("INSERT INTO users(username,password,role,emp_id) VALUES(?,?,?,?)",
                      (request.form["username"], generate_password_hash(request.form["password"]), request.form["role"], request.form.get("emp_id") or None))
            d.commit(); flash("User created.")
        except Exception:
            d.rollback(); flash("This username already exists.")
        return redirect("/users")
    rows = d.execute("SELECT id,username,role FROM users").fetchall()
    return page("""<h2>Users</h2><div class="card"><form class="row" method="post"><input name="username" placeholder="Username" required>
    <input name="password" type="password" placeholder="Password" required><select name="role"><option value="hr">HR</option><option value="admin">Admin</option><option value="employee">Employee</option></select><input name="emp_id" type="number" placeholder="Employee ID (for employee role)">
    <button>Create user</button></form><table class="ft" data-ph="Search users..."><tr><th>ID</th><th>Username</th><th>Role</th><th></th></tr>
    {% for r in rows %}<tr><td>{{r.id}}</td><td>{{r.username}}</td><td>{{r.role}}</td><td class="noprint"><a href="/edit/users/{{r.id}}">Edit</a> <form method="post" action="/delete/users/{{r.id}}" style="display:inline" onsubmit="return confirm('Are you sure you want to delete this?')"><button>Delete</button></form></td></tr>{% endfor %}</table></div>""", rows=rows)


# =====================================================================
# FEATURES: 1) Payroll + Salary slip PDF  2) Excel export  3) Employee login
# (4) PostgreSQL: enabled through DATABASE_URL above)
# =====================================================================
WORK_DAYS = 26      # working days in a month
LATE_PER_DAY = 3    # every 3 Late marks = 1 day deduction

OT_HOURS_PER_DAY = 8   # hours in a working day (to get the hourly rate)
OT_RATE = 2.0          # overtime double rate

def loan_due(l, month):
    """Installment due this month and the amount remaining after it. Installments are deducted automatically every month."""
    gone = (int(month[:4]) - int(l["start_month"][:4])) * 12 + int(month[5:7]) - int(l["start_month"][5:7])
    if gone < 0: return 0.0, float(l["amount"])
    remaining = max(float(l["amount"]) - float(l["installment"]) * gone, 0)
    due = min(float(l["installment"]), remaining)
    return due, remaining - due

def calc(d, e, month):
    rows = d.execute("SELECT status, COUNT(*) AS n FROM attendance WHERE emp_id=? AND day LIKE ? GROUP BY status",
                     (e["id"], month + "%")).fetchall()
    c = {r["status"]: r["n"] for r in rows}
    ded_days = c.get("Absent", 0) + c.get("Late", 0) // LATE_PER_DAY + 0.5 * c.get("Half Day", 0)
    sal = float(e["salary"] or 0); ded = round(sal / WORK_DAYS * ded_days)
    allow = sum(float(e[k] or 0) for k in ("allowances", "house_rent", "medical_allowance", "bonus")); tax = float(e["tax_deduction"] or 0)
    hrs = float(d.execute("SELECT COALESCE(SUM(hours),0) FROM overtime WHERE emp_id=? AND day LIKE ?", (e["id"], month + "%")).fetchone()[0])
    ot_pay = round((float(e["ot_rate"] or 0) or sal / WORK_DAYS / OT_HOURS_PER_DAY * OT_RATE) * hrs)
    loan = round(sum(loan_due(l, month)[0] for l in d.execute("SELECT * FROM loans WHERE emp_id=?", (e["id"],)).fetchall()))
    return dict(present=c.get("Present", 0), absent=c.get("Absent", 0), late=c.get("Late", 0), leave=c.get("Leave", 0), half=c.get("Half Day", 0),
                ded_days=ded_days, deduction=ded, ot_hours=hrs, ot_pay=ot_pay, loan_ded=loan, allow=allow, tax=tax, net=sal + allow - ded + ot_pay - loan - tax)

@app.route("/payroll")
@login_required("admin", "hr")
def payroll():
    d = db(); m = request.args.get("month", date.today().isoformat()[:7])
    rows = [(e, calc(d, e, m)) for e in d.execute("SELECT * FROM employees WHERE active=1 AND COALESCE(status,'Active') NOT IN ('Resigned','Terminated')").fetchall()]
    return page("""<h2>Payroll</h2><div class="card"><form class="row"><input type="month" name="month" value="{{m}}"><button>Show</button></form>
    <p>Rule: every Absent = 1 day deduction, every {{L}} Late = 1 day. {{W}} days in a month. Half Day = 0.5 day. Overtime = the employee's own hourly rate, or (basic / {{W}} / 8) x {{O}} per hour when none is set.</p>
    <table class="ft" data-ph="Search employee..."><tr><th>Employee</th><th>Basic</th><th>Allowances</th><th>Present</th><th>Absent</th><th>Half</th><th>Late</th><th>Absent/Late ded.</th><th>OT hrs</th><th>OT pay</th><th>Advance/Loan</th><th>Tax/Ded.</th><th>Net pay</th><th>Slip</th></tr>
    {% for e,c in rows %}<tr><td>{{e.name}}</td><td>{{'{:,.0f}'.format(e.salary or 0)}}</td><td>{{'{:,.0f}'.format(c.allow)}}</td><td>{{c.present}}</td><td>{{c.absent}}</td><td>{{c.half}}</td><td>{{c.late}}</td>
    <td>{{'{:,.0f}'.format(c.deduction)}}</td><td>{{c.ot_hours}}</td><td>{{'{:,.0f}'.format(c.ot_pay)}}</td><td>{{'{:,.0f}'.format(c.loan_ded)}}</td><td>{{'{:,.0f}'.format(c.tax)}}</td><td><b>{{'{:,.0f}'.format(c.net)}}</b></td>
    <td><a href="/payslip/{{e.id}}/{{m}}.pdf">PDF</a></td></tr>{% endfor %}</table>
    <h3>Total payable: {{'{:,.0f}'.format(total)}}</h3></div>""", rows=rows, m=m, L=LATE_PER_DAY, W=WORK_DAYS, O=OT_RATE,
                total=sum(c["net"] for _, c in rows))

@app.route("/payslip/<int:eid>/<month>.pdf")
@login_required()
def payslip_pdf(eid, month):
    if session["role"] not in ("admin", "hr") and session.get("emp_id") != eid:
        return "Access denied", 403
    d = db(); e = d.execute("SELECT * FROM employees WHERE id=?", (eid,)).fetchone()
    if not e: return "Employee not found", 404
    c = calc(d, e, month)
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    buf = io.BytesIO(); p = canvas.Canvas(buf, pagesize=A4); y = 800
    p.setFont("Helvetica-Bold", 18); p.drawString(50, y, "ERP Pro - Salary Slip"); y -= 24
    p.setFont("Helvetica", 12); p.drawString(50, y, "Month: " + month); y -= 30
    for k, v in [("Name", e["name"]), ("Father's name", e["father_name"]), ("CNIC", e["cnic"]), ("Department", e["department"]), ("Designation", e["designation"]),
                 ("Basic salary", f"{float(e['salary'] or 0):,.0f}"), ("Allowances / bonus", f"+ {c['allow']:,.0f}"), ("Present days", c["present"]), ("Absent days", c["absent"]), ("Half days", c["half"]),
                 ("Late days", c["late"]), ("Leave days", c["leave"]),
                 (f"Deduction ({c['ded_days']} days)", f"- {c['deduction']:,.0f}"),
                 (f"Overtime ({c['ot_hours']:g} hrs)", f"+ {c['ot_pay']:,.0f}"), ("Advance/Loan installment", f"- {c['loan_ded']:,.0f}"), ("Tax / deductions", f"- {c['tax']:,.0f}"),
                 ("Payment method", e["pay_method"]), ("Bank / account", f"{e['bank_name'] or ''} {e['account_no'] or ''}".strip())]:
        p.drawString(50, y, str(k)); p.drawString(250, y, str(v or "")); y -= 20
    p.setFont("Helvetica-Bold", 14); y -= 10
    p.drawString(50, y, "NET PAY"); p.drawString(250, y, f"{c['net']:,.0f}")
    p.showPage(); p.save(); buf.seek(0)
    return send_file(buf, mimetype="application/pdf", download_name=f"salary_slip_{e['name']}_{month}.pdf")

# ---------- Leave quota, balance, holidays ----------
LEAVE_QUOTA = {"Casual": 10, "Sick": 8, "Annual": 14}   # days per year, change as needed

def holiday_set(d):
    return {r["day"] for r in d.execute("SELECT day FROM holidays").fetchall()}

def leave_days(d, l, year=None):
    """Days counted for a leave request (weekly off and public holidays are not counted)."""
    hol = holiday_set(d); a, b = date.fromisoformat(l["start_date"]), date.fromisoformat(l["end_date"]); n = 0
    while a <= b:
        if a.weekday() not in WEEKLY_OFF and a.isoformat() not in hol and (year is None or a.year == year): n += 1
        a += timedelta(days=1)
    return n

def leave_balance(d, eid, year=None):
    year = year or date.today().year; used = {k: 0 for k in LEAVE_QUOTA}
    for l in d.execute("SELECT * FROM leaves WHERE emp_id=? AND status='Approved'", (eid,)).fetchall():
        k = l["leave_type"] if l["leave_type"] in LEAVE_QUOTA else "Casual"; used[k] += leave_days(d, l, year)
    return {k: dict(quota=q, used=used[k], left=q - used[k]) for k, q in LEAVE_QUOTA.items()}

# ---------- Employee self-service portal ----------
@app.route("/me")
@login_required("employee")
def me():
    d = db(); eid = session.get("emp_id"); m = date.today().isoformat()[:7]
    e = d.execute("SELECT * FROM employees WHERE id=?", (eid,)).fetchone()
    if not e: return page("<div class='card'>Your account is not linked to an employee record. Please contact the Admin.</div>")
    att = d.execute("SELECT * FROM attendance WHERE emp_id=? AND day LIKE ? ORDER BY day DESC", (eid, m + "%")).fetchall()
    lv = d.execute("SELECT * FROM leaves WHERE emp_id=? ORDER BY id DESC", (eid,)).fetchall()
    cr = d.execute("SELECT * FROM corrections WHERE emp_id=? ORDER BY id DESC", (eid,)).fetchall()
    has_photo = bool(d.execute("SELECT 1 FROM documents WHERE emp_id=? AND kind='photo'", (eid,)).fetchone())
    return page("""<div class="pgh"><span class="av2">{% if has_photo %}<img src="/emp/{{e.id}}/photo" alt="">{% else %}{{(e.name or '?')[:1]|upper}}{% endif %}</span>
    <div><h2>Welcome, {{e.name}}</h2><p>{{e.designation}}{% if e.department %} &middot; {{e.department}}{% endif %} &middot; Staff code {{e.staff_code}}</p></div></div>
    <div class="stats">{% for k,v in bal.items() %}<div class="sc"><span class="ico b">{{icon('calendar')}}</span><div><small>{{k}} leave balance</small><strong>{{v.left}} <em>/ {{v.quota}} days</em></strong></div></div>{% endfor %}</div>
    <div class="card"><h3>My attendance ({{m}})</h3><table><tr><th>Date</th><th>Status</th><th>Check-in</th><th>Check-out</th><th>Late (min)</th><th>Left early (min)</th></tr>
    {% for a in att %}<tr><td>{{a.day}}</td><td>{{a.status}}</td><td>{{a.check_in or ''}}</td><td>{{a.check_out or ''}}</td><td>{{a.late_min or ''}}</td><td>{{a.early_min or ''}}</td></tr>{% endfor %}</table>
    <p><a href="/payslip/{{e.id}}/{{m}}.pdf">This month's salary slip (PDF)</a></p></div>
    <div class="card"><h3>Request leave</h3><form class="row" method="post" action="/leave/apply"><select name="leave_type">{% for k in bal %}<option>{{k}}</option>{% endfor %}</select>
    <input type="date" name="start" required><input type="date" name="end" required><input name="reason" placeholder="Reason" required><button>Send request</button></form>
    <table><tr><th>Type</th><th>From</th><th>To</th><th>Reason</th><th>Status</th></tr>{% for l in lv %}<tr><td>{{l.leave_type or 'Casual'}}</td><td>{{l.start_date}}</td><td>{{l.end_date}}</td><td>{{l.reason}}</td><td>{{l.status}}</td></tr>{% endfor %}</table></div>
    <div class="card"><h3>Attendance correction request</h3><p style="color:var(--mut)">Marked wrongly? Send the correct details and HR will review.</p>
    <form class="row" method="post" action="/correction/apply"><input type="date" name="day" required><select name="new_status">{% for s in ['Present','Late','Half Day','Absent','Leave'] %}<option>{{s}}</option>{% endfor %}</select>
    <input type="time" name="check_in" title="Check-in"><input type="time" name="check_out" title="Check-out"><input name="reason" placeholder="Reason" required><button>Send correction</button></form>
    <table><tr><th>Date</th><th>Requested</th><th>In</th><th>Out</th><th>Reason</th><th>Status</th></tr>{% for c in cr %}<tr><td>{{c.day}}</td><td>{{c.new_status}}</td><td>{{c.check_in or ''}}</td><td>{{c.check_out or ''}}</td><td>{{c.reason}}</td><td>{{c.state}}</td></tr>{% endfor %}</table></div>""",
                e=e, att=att, m=m, lv=lv, cr=cr, bal=leave_balance(d, eid), has_photo=has_photo)

# ---------- Attendance correction requests ----------
@app.route("/correction/apply", methods=["POST"])
@login_required("employee")
def correction_apply():
    f = request.form
    try: day = date.fromisoformat(f["day"]).isoformat()
    except ValueError: flash("Please enter a valid date."); return redirect("/me")
    st = f.get("new_status") if f.get("new_status") in STATUSES else "Present"
    db().execute("INSERT INTO corrections(emp_id,day,new_status,check_in,check_out,reason,state,created) VALUES(?,?,?,?,?,?,'Pending',?)",
                 (session.get("emp_id"), day, st, clean_time(f.get("check_in")), clean_time(f.get("check_out")), f.get("reason", ""), date.today().isoformat()))
    db().commit(); flash("Correction request sent to HR."); return redirect("/me")

@app.route("/corrections")
@login_required("admin", "hr")
def corrections():
    rows = db().execute("SELECT c.*, e.name FROM corrections c JOIN employees e ON e.id=c.emp_id ORDER BY (c.state='Pending') DESC, c.id DESC").fetchall()
    return page("""<h2>Attendance correction requests</h2><div class="card"><p>Approving applies the requested status and times to that day's attendance.</p>
    <table class="ft" data-ph="Search requests..."><tr><th>Employee</th><th>Date</th><th>Requested</th><th>In</th><th>Out</th><th>Reason</th><th>Status</th><th></th></tr>
    {% for r in rows %}<tr><td>{{r.name}}</td><td>{{r.day}}</td><td>{{r.new_status}}</td><td>{{r.check_in or ''}}</td><td>{{r.check_out or ''}}</td><td>{{r.reason}}</td><td>{{r.state}}</td><td class="noprint">
    {% if r.state=='Pending' %}<form method="post" action="/corrections/{{r.id}}/approve" style="display:inline"><button>Approve</button></form>
    <form method="post" action="/corrections/{{r.id}}/reject" style="display:inline"><button>Reject</button></form>{% endif %}</td></tr>{% endfor %}</table></div>""", rows=rows)

@app.route("/corrections/<int:i>/<action>", methods=["POST"])
@login_required("admin", "hr")
def correction_action(i, action):
    d = db(); c = d.execute("SELECT * FROM corrections WHERE id=? AND state='Pending'", (i,)).fetchone()
    if not c or action not in ("approve", "reject"): return redirect("/corrections")
    if action == "approve":
        st, cin, cout = c["new_status"], c["check_in"], c["check_out"]; late = early = None
        if st in ("Present", "Late", "Half Day"):
            if cin: late = max(0, (tmin(cin) or 0) - tmin(SHIFT_START)) or None
            if cout and st != "Half Day": early = max(0, tmin(SHIFT_END) - (tmin(cout) or 0)) or None
            if st == "Present" and late and late > GRACE_MIN: st = "Late"
        else: cin = cout = None
        d.execute("""INSERT INTO attendance(emp_id,day,status,check_in,check_out,late_min,early_min) VALUES(?,?,?,?,?,?,?) ON CONFLICT(emp_id,day) DO UPDATE SET
            status=excluded.status, check_in=excluded.check_in, check_out=excluded.check_out, late_min=excluded.late_min, early_min=excluded.early_min""", (c["emp_id"], c["day"], st, cin, cout, late, early))
    d.execute("UPDATE corrections SET state=? WHERE id=?", ("Approved" if action == "approve" else "Rejected", i)); d.commit()
    flash("Correction " + ("approved and applied." if action == "approve" else "rejected.")); return redirect("/corrections")

# ---------- Holidays & weekly off ----------
@app.route("/holidays", methods=["GET", "POST"])
@login_required("admin", "hr")
def holidays():
    d = db()
    if request.method == "POST":
        d.execute("INSERT INTO holidays(day,name) VALUES(?,?)", (request.form["day"], request.form["name"])); d.commit(); flash("Holiday added."); return redirect("/holidays")
    rows = d.execute("SELECT * FROM holidays ORDER BY day").fetchall()
    return page("""<h2>Holidays &amp; weekly off</h2><div class="card"><p>Weekly off: <b>{{off}}</b> (change <code>WEEKLY_OFF</code> in app.py). Holidays are shaded in the attendance chart and are not counted as leave days.</p>
    <form class="row" method="post"><input type="date" name="day" required><input name="name" placeholder="Holiday name (e.g. Eid ul Fitr)" required><button>Add holiday</button></form>
    <table class="ft" data-ph="Search holidays..."><tr><th>Date</th><th>Day</th><th>Holiday</th><th></th></tr>
    {% for r in rows %}<tr><td>{{r.day}}</td><td>{{wd(r.day)}}</td><td>{{r.name}}</td>""" + act("holidays") + """</tr>{% endfor %}</table></div>""",
                rows=rows, off=", ".join(["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"][i] for i in WEEKLY_OFF),
                wd=lambda x: date.fromisoformat(x).strftime("%A"))

# ---------- Change password (every user) ----------
@app.route("/password", methods=["GET", "POST"])
@login_required()
def change_password():
    if request.method == "POST":
        d = db(); f = request.form
        u = d.execute("SELECT * FROM users WHERE username=?", (session["user"],)).fetchone()
        if not check_password_hash(u["password"], f["old"]): flash("Current password is incorrect.")
        elif len(f["new"]) < 6: flash("New password must be at least 6 characters.")
        elif f["new"] != f["confirm"]: flash("New passwords do not match.")
        else:
            d.execute("UPDATE users SET password=? WHERE id=?", (generate_password_hash(f["new"]), u["id"])); d.commit()
            flash("Password updated."); return redirect("/")
        return redirect("/password")
    return page("""<h2>Change password</h2><div class="card" style="max-width:380px"><form method="post">
    <p><input type="password" name="old" placeholder="Current password" style="width:100%" required></p>
    <p><input type="password" name="new" placeholder="New password" style="width:100%" required></p>
    <p><input type="password" name="confirm" placeholder="Confirm new password" style="width:100%" required></p><button>Save password</button></form></div>""")

# ---------- Monthly attendance report ----------
SYMBOL = {"Present": "P", "Absent": "A", "Late": "L", "Leave": "V", "Half Day": "H"}

def month_report(d, m):
    y, mo = int(m[:4]), int(m[5:7]); days = calendar.monthrange(y, mo)[1]
    emps = d.execute("SELECT id, name FROM employees WHERE active=1 AND COALESCE(status,'Active') NOT IN ('Resigned','Terminated') ORDER BY id").fetchall()
    got = {(r["emp_id"], r["day"]): r["status"] for r in d.execute("SELECT emp_id, day, status FROM attendance WHERE day LIKE ?", (m + "%",)).fetchall()}
    rows = []
    for e in emps:
        cells = [SYMBOL.get(got.get((e["id"], f"{m}-{n:02d}")), "") for n in range(1, days + 1)]
        rows.append((e, cells, [cells.count(x) for x in "PALV"]))
    return days, rows

@app.route("/attendance/report")
@login_required("admin", "hr")
def attendance_report():   # the monthly chart now lives on the Attendance tab
    return redirect("/attendance?month=" + request.args.get("month", ""))

@app.route("/attendance/report.xlsx")
@login_required("admin", "hr")
def attendance_report_xlsx():
    from openpyxl import Workbook
    m = request.args.get("month", date.today().isoformat()[:7]); days, rows = month_report(db(), m)
    wb = Workbook(); ws = wb.active; ws.title = m
    ws.append(["ID", "Name"] + list(range(1, days + 1)) + ["Present", "Absent", "Late", "Leave"])
    for e, cells, t in rows: ws.append([e["id"], e["name"]] + cells + t)
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return send_file(buf, download_name=f"attendance_{m}.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

# ---------- Leave management ----------
@app.route("/leave/apply", methods=["POST"])
@login_required("employee")
def leave_apply():
    f = request.form; eid = session.get("emp_id")
    try:
        a, b = date.fromisoformat(f["start"]), date.fromisoformat(f["end"])
        if b < a or (b - a).days > 30: raise ValueError
    except ValueError:
        flash("Please enter valid dates (the end date cannot be before the start date, and the range cannot exceed 31 days)."); return redirect("/me")
    lt = f.get("leave_type") if f.get("leave_type") in LEAVE_QUOTA else "Casual"
    db().execute("INSERT INTO leaves(emp_id,start_date,end_date,reason,status,leave_type) VALUES(?,?,?,?,'Pending',?)", (eid, a.isoformat(), b.isoformat(), f["reason"], lt))
    db().commit(); flash("Leave request sent to HR."); return redirect("/me")

@app.route("/leaves")
@login_required("admin", "hr")
def leaves():
    rows = []; d = db()
    for l in d.execute("SELECT l.*, e.name FROM leaves l JOIN employees e ON e.id=l.emp_id ORDER BY (l.status='Pending') DESC, l.id DESC").fetchall():
        lt = l["leave_type"] or "Casual"; bal = leave_balance(d, l["emp_id"]).get(lt)
        rows.append(dict(l, days=leave_days(d, l), lt=lt, left=bal["left"] if bal else ""))
    return page("""<h2>Leave requests</h2><div class="card"><p>Approving a request marks attendance as "Leave" for those dates (weekly off and holidays excluded), and no payroll deduction is made. Balance = yearly quota minus approved days of that type.</p>
    <table class="ft" data-ph="Search leave requests..."><tr><th>Employee</th><th>Type</th><th>Balance left</th><th>From</th><th>To</th><th>Days</th><th>Reason</th><th>Status</th><th></th></tr>
    {% for r in rows %}<tr><td>{{r.name}}</td><td>{{r.lt}}</td><td>{{r.left}} day(s)</td><td>{{r.start_date}}</td><td>{{r.end_date}}</td><td>{{r.days}}</td><td>{{r.reason}}</td><td>{{r.status}}</td><td class="noprint">
    {% if r.status=='Pending' %}<form method="post" action="/leaves/{{r.id}}/approve" style="display:inline"><button>Approve</button></form>
    <form method="post" action="/leaves/{{r.id}}/reject" style="display:inline"><button>Reject</button></form>{% endif %}</td></tr>{% endfor %}</table></div>""", rows=rows)

@app.route("/leaves/<int:i>/<action>", methods=["POST"])
@login_required("admin", "hr")
def leave_action(i, action):
    d = db(); l = d.execute("SELECT * FROM leaves WHERE id=? AND status='Pending'", (i,)).fetchone()
    if not l or action not in ("approve", "reject"): return redirect("/leaves")
    if action == "approve":
        a, b = date.fromisoformat(l["start_date"]), date.fromisoformat(l["end_date"])
        while a <= b:
            if a.weekday() not in WEEKLY_OFF and a.isoformat() not in holiday_set(d):   # weekly off and public holidays are skipped
                cur = d.execute("SELECT status FROM attendance WHERE emp_id=? AND day=?", (l["emp_id"], a.isoformat())).fetchone()
                if not cur or cur["status"] not in ("Present", "Late"):
                    d.execute("INSERT INTO attendance(emp_id,day,status) VALUES(?,?,'Leave') ON CONFLICT(emp_id,day) DO UPDATE SET status=excluded.status", (l["emp_id"], a.isoformat()))
            a += timedelta(days=1)
    d.execute("UPDATE leaves SET status=? WHERE id=?", ("Approved" if action == "approve" else "Rejected", i)); d.commit()
    flash("Leave request " + ("approved" if action == "approve" else "rejected") + "."); return redirect("/leaves")

# ---------- Overtime ----------
def act(t):
    return ('<td class="noprint"><a href="/edit/%s/{{r.id}}">Edit</a> <form method="post" action="/delete/%s/{{r.id}}" style="display:inline" '
            'onsubmit="return confirm(\'Are you sure you want to delete this?\')"><button>Delete</button></form></td>') % (t, t)

@app.route("/overtime", methods=["GET", "POST"])
@login_required("admin", "hr")
def overtime():
    d = db()
    if request.method == "POST":
        f = request.form
        d.execute("INSERT INTO overtime(emp_id,day,hours) VALUES(?,?,?)", (f["emp"], f["day"], f["hours"])); d.commit()
        flash("Overtime added."); return redirect("/overtime")
    m = request.args.get("month", date.today().isoformat()[:7])
    rows = d.execute("SELECT o.*, e.name FROM overtime o JOIN employees e ON e.id=o.emp_id WHERE o.day LIKE ? ORDER BY o.day DESC", (m + "%",)).fetchall()
    return page("""<h2>Overtime</h2><div class="card"><form class="row" method="post"><select name="emp" required>
    {% for e in emps %}<option value="{{e.id}}">{{e.id}} - {{e.name}}</option>{% endfor %}</select>
    <input type="date" name="day" value="{{today}}"><input name="hours" type="number" step="0.5" min="0.5" placeholder="Hours" required><button>Add overtime</button></form>
    <form class="row"><input type="month" name="month" value="{{m}}"><button>Show</button></form>
    <table class="ft" data-ph="Search overtime..."><tr><th>Employee</th><th>Date</th><th>Hours</th><th></th></tr>{% for r in rows %}<tr><td>{{r.name}}</td><td>{{r.day}}</td><td>{{r.hours}}</td>""" + act("overtime") + """</tr>{% endfor %}</table>
    <p>Total hours: <b>{{total}}</b></p></div>""", emps=d.execute("SELECT id, name FROM employees WHERE active=1 AND COALESCE(status,'Active') NOT IN ('Resigned','Terminated') ORDER BY id").fetchall(),
                rows=rows, m=m, today=date.today().isoformat(), total=sum(r["hours"] for r in rows))

# ---------- Advance / Loan ----------
@app.route("/loans", methods=["GET", "POST"])
@login_required("admin", "hr")
def loans():
    d = db(); now = date.today().isoformat()[:7]
    if request.method == "POST":
        f = request.form
        if float(f["installment"]) <= 0 or float(f["amount"]) <= 0:
            flash("Amount and installment must be greater than 0."); return redirect("/loans")
        d.execute("INSERT INTO loans(emp_id,amount,installment,start_month,note) VALUES(?,?,?,?,?)",
                  (f["emp"], f["amount"], f["installment"], f["start"], f["note"])); d.commit()
        flash("Advance/loan added. The installment will be deducted from payroll automatically."); return redirect("/loans")
    rows = [dict(l, left=loan_due(l, now)[1]) for l in d.execute(
        "SELECT l.*, e.name FROM loans l JOIN employees e ON e.id=l.emp_id ORDER BY l.id DESC").fetchall()]
    return page("""<h2>Advance / Loan</h2><div class="card"><form class="row" method="post"><select name="emp" required>
    {% for e in emps %}<option value="{{e.id}}">{{e.id}} - {{e.name}}</option>{% endfor %}</select>
    <input name="amount" type="number" step="any" placeholder="Total amount" required><input name="installment" type="number" step="any" placeholder="Monthly installment" required>
    <input type="month" name="start" value="{{now}}" required><input name="note" placeholder="Note"><button>Add</button></form>
    <p>The installment is deducted from payroll automatically from the start month until the amount is fully repaid. "Remaining" is as of the current month.</p>
    <table class="ft" data-ph="Search loans..."><tr><th>Employee</th><th>Total</th><th>Installment</th><th>Start</th><th>Remaining</th><th>Note</th><th></th></tr>
    {% for r in rows %}<tr><td>{{r.name}}</td><td>{{'{:,.0f}'.format(r.amount)}}</td><td>{{'{:,.0f}'.format(r.installment)}}</td><td>{{r.start_month}}</td>
    <td class="{{'low' if r.left>0}}">{{'{:,.0f}'.format(r.left)}}{{' (cleared)' if r.left==0}}</td><td>{{r.note}}</td>""" + act("loans") + """</tr>{% endfor %}</table></div>""",
                emps=d.execute("SELECT id, name FROM employees WHERE active=1 AND COALESCE(status,'Active') NOT IN ('Resigned','Terminated') ORDER BY id").fetchall(), rows=rows, now=now)

# ---------- Selected employees: attendance report / information as PDF, Excel, Word ----------
MIMES = {"pdf": "application/pdf", "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
         "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document"}
FILL = {"Present": "E3F6EC", "Absent": "FDE8E8", "Late": "FFF1D6", "Leave": "E8EFFF", "Half Day": "EFE8FF"}

def att_data(d, ids, m):
    y, mo = int(m[:4]), int(m[5:7]); qs = ",".join("?" * len(ids))
    emps = d.execute(f"SELECT * FROM employees WHERE active=1 AND id IN ({qs}) ORDER BY id", ids).fetchall()
    recs = {(r["emp_id"], r["day"]): r for r in d.execute(f"SELECT * FROM attendance WHERE day LIKE ? AND emp_id IN ({qs})", [m + "%"] + ids).fetchall()}
    out = []
    for e in emps:
        days = []
        for i in range(1, calendar.monthrange(y, mo)[1] + 1):
            iso = f"{m}-{i:02d}"; r = recs.get((e["id"], iso)); cin = (r["check_in"] or "") if r else ""; cout = (r["check_out"] or "") if r else ""
            days.append(dict(day=iso, wd=WD[date(y, mo, i).weekday()], status=r["status"] if r else "", cin=cin, cout=cout,
                             late=(r["late_min"] or 0) if r else 0, hours=worked(cin, cout)))
        out.append(dict(e=e, days=days, tot={k: sum(1 for x in days if x["status"] == k) for k in STATUSES},
                        late_min=sum(x["late"] for x in days), hours=round(sum(x["hours"] for x in days), 1)))
    return out

def info_rows(e):
    clean = lambda l: re.sub(r"\s*\(.*?\)", "", l)
    return [("Employee ID", e["id"]), ("Staff code", e["staff_code"])] + [(clean(l), e[k] if e[k] is not None else "") for k, l, _ in EMP_FIELDS]

def sum_row(x):
    t = x["tot"]; return [x["e"]["id"], x["e"]["name"], x["e"]["department"] or "", t["Present"], t["Absent"], t["Late"], t["Leave"], t["Half Day"], x["late_min"], x["hours"]]
SUM_HEAD = ["ID", "Employee", "Department", "Present", "Absent", "Late", "Leave", "Half day", "Late min", "Hours"]
DAY_HEAD = ["Date", "Day", "Status", "Check-in", "Check-out", "Late (min)", "Hours"]
day_row = lambda z: [z["day"], z["wd"], z["status"] or "-", z["cin"] or "-", z["cout"] or "-", z["late"] or "-", z["hours"] or "-"]
month_label = lambda m: date(int(m[:4]), int(m[5:7]), 1).strftime("%B %Y")

def pdf_doc(build):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, PageBreak, Spacer
    buf = io.BytesIO(); doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=30, rightMargin=30, topMargin=34, bottomMargin=34); ss = getSampleStyleSheet()
    def tbl(rows, widths=None, head=True):
        t = Table(rows, repeatRows=1 if head else 0, colWidths=widths)
        st = [("FONTSIZE", (0, 0), (-1, -1), 8.5), ("GRID", (0, 0), (-1, -1), .25, colors.HexColor("#d5dbe6")), ("TOPPADDING", (0, 0), (-1, -1), 2.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5)]
        if head: st += [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0b1628")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f5f7fb")])]
        t.setStyle(TableStyle(st)); return t
    doc.build(build(Paragraph, Spacer, PageBreak, ss, tbl, colors, TableStyle)); buf.seek(0); return buf

def att_pdf(data, m, detail):
    def build(P, Sp, PB, ss, tbl, colors, TS):
        el = [P("ERP Pro - Monthly Attendance Report", ss["Title"]), P(f"{month_label(m)} &nbsp;|&nbsp; {len(data)} employee(s)", ss["Normal"]), Sp(1, 12), tbl([SUM_HEAD] + [sum_row(x) for x in data])]
        if detail:
            for x in data:
                t = x["tot"]; el += [PB(), P(f'{x["e"]["name"]} (ID {x["e"]["id"]}) - {month_label(m)}', ss["Heading2"]),
                    P(f'Present {t["Present"]} | Absent {t["Absent"]} | Late {t["Late"]} ({x["late_min"]} min) | Leave {t["Leave"]} | Half day {t["Half Day"]} | Hours worked {x["hours"]}', ss["Normal"]), Sp(1, 8)]
                tb = tbl([DAY_HEAD] + [day_row(z) for z in x["days"]])
                for i, z in enumerate(x["days"], 1):
                    if z["status"] in FILL: tb.setStyle(TS([("BACKGROUND", (2, i), (2, i), colors.HexColor("#" + FILL[z["status"]]))]))
                el.append(tb)
        return el
    return pdf_doc(build)

def info_pdf(emps, m, detail):
    def build(P, Sp, PB, ss, tbl, colors, TS):
        el = []
        for n, e in enumerate(emps):
            if n: el.append(PB())
            t = tbl([[a, "" if b is None else str(b)[:160]] for a, b in info_rows(e)], widths=[150, 370], head=False)
            t.setStyle(TS([("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#f1f4fa")), ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, -1), 10), ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]))
            el += [P("ERP Pro - Employee Information Sheet", ss["Title"]), P(f'<b>{e["name"]}</b> &nbsp;|&nbsp; {e["designation"] or ""} &nbsp;|&nbsp; {e["department"] or ""}', ss["Normal"]), Sp(1, 10)]
            pp = photo_path(e["id"])
            if pp:
                try:
                    from reportlab.platypus import Image
                    el += [Image(pp, width=80, height=80, kind="proportional", hAlign="LEFT"), Sp(1, 8)]
                except Exception: pass
            el.append(t)
        return el
    return pdf_doc(build)

def att_xlsx(data, m, detail):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook(); ws = wb.active; ws.title = "Chart"; nd = len(data[0]["days"]) if data else 0
    ws.append(["ID", "Employee"] + list(range(1, nd + 1)) + ["Present", "Absent", "Late", "Leave", "Half day", "Late min", "Hours"])
    for x in data:
        ws.append([x["e"]["id"], x["e"]["name"]] + [SYMBOL.get(z["status"], "") for z in x["days"]] + [x["tot"][k] for k in STATUSES] + [x["late_min"], x["hours"]])
        for i, z in enumerate(x["days"]):
            if z["status"]: c = ws.cell(ws.max_row, 3 + i); c.fill = PatternFill("solid", fgColor=FILL[z["status"]]); c.alignment = Alignment(horizontal="center")
    ws2 = wb.create_sheet("Summary"); ws2.append(SUM_HEAD); [ws2.append(sum_row(x)) for x in data]
    if detail:
        for x in data:
            w = wb.create_sheet(re.sub(r"[\[\]:*?/\\]", "", f'{x["e"]["id"]} {x["e"]["name"]}')[:31]); w.append(DAY_HEAD)
            for z in x["days"]:
                w.append(day_row(z))
                if z["status"]: w.cell(w.max_row, 3).fill = PatternFill("solid", fgColor=FILL[z["status"]])
            w.column_dimensions["A"].width = 13
    for sh in wb.worksheets:
        for c in sh[1]: c.font = Font(bold=True)
    ws.column_dimensions["B"].width = 26; ws2.column_dimensions["B"].width = 26
    buf = io.BytesIO(); wb.save(buf); buf.seek(0); return buf

def info_xlsx(emps, m, detail):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook(); ws = wb.active; ws.title = "Employees"; ws.append([a for a, _ in info_rows(emps[0])] if emps else [])
    for e in emps: ws.append([b for _, b in info_rows(e)])
    for c in ws[1]: c.font = Font(bold=True)
    buf = io.BytesIO(); wb.save(buf); buf.seek(0); return buf

def docx_table(doc, rows, head=True):
    t = doc.add_table(rows=0, cols=len(rows[0])); t.style = "Table Grid"
    for i, r in enumerate(rows):
        cells = t.add_row().cells
        for c, v in zip(cells, r):
            c.text = "" if v is None else str(v)
            for p in c.paragraphs:
                for run in p.runs:
                    run.font.size = __import__("docx").shared.Pt(8.5)
                    if head and i == 0: run.bold = True
    return t

def att_docx(data, m, detail):
    from docx import Document
    doc = Document(); doc.add_heading("ERP Pro - Monthly Attendance Report", 0); doc.add_paragraph(f"{month_label(m)} | {len(data)} employee(s)")
    docx_table(doc, [SUM_HEAD] + [sum_row(x) for x in data])
    if detail:
        for x in data:
            doc.add_page_break(); doc.add_heading(f'{x["e"]["name"]} (ID {x["e"]["id"]}) - {month_label(m)}', 2)
            t = x["tot"]; doc.add_paragraph(f'Present {t["Present"]} | Absent {t["Absent"]} | Late {t["Late"]} ({x["late_min"]} min) | Leave {t["Leave"]} | Half day {t["Half Day"]} | Hours worked {x["hours"]}')
            docx_table(doc, [DAY_HEAD] + [day_row(z) for z in x["days"]])
    buf = io.BytesIO(); doc.save(buf); buf.seek(0); return buf

def info_docx(emps, m, detail):
    from docx import Document
    doc = Document()
    for n, e in enumerate(emps):
        if n: doc.add_page_break()
        doc.add_heading("ERP Pro - Employee Information Sheet", 1); doc.add_paragraph(e["name"]); docx_table(doc, info_rows(e), head=False)
    buf = io.BytesIO(); doc.save(buf); buf.seek(0); return buf

@app.route("/bulk/download")
@login_required("admin", "hr")
def bulk_download():
    ids = [int(x) for x in request.args.get("ids", "").split(",") if x.strip().isdigit()][:500]
    kind, fmt, m = request.args.get("kind", ""), request.args.get("fmt", ""), request.args.get("month", "")
    if not ids or kind not in ("att", "info") or fmt not in MIMES: return "Bad request", 400
    if not re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", m): m = date.today().isoformat()[:7]
    d = db(); detail = request.args.get("detail", "1") == "1"
    if kind == "att": data = att_data(d, ids, m)
    else: data = d.execute(f"SELECT * FROM employees WHERE active=1 AND id IN ({','.join('?' * len(ids))}) ORDER BY id", ids).fetchall()
    if not data: return "No matching employees", 404
    fn = {("att", "pdf"): att_pdf, ("att", "xlsx"): att_xlsx, ("att", "docx"): att_docx, ("info", "pdf"): info_pdf, ("info", "xlsx"): info_xlsx, ("info", "docx"): info_docx}[(kind, fmt)]
    name = (f"attendance_{m}_{len(data)}_employees" if kind == "att" else f"employee_information_{len(data)}") + "." + fmt
    return send_file(fn(data, m, detail), mimetype=MIMES[fmt], download_name=name, as_attachment=True)

@app.route("/employees/bulk_delete", methods=["POST"])
@login_required("admin", "hr")
def employees_bulk_delete():
    ids = [int(x) for x in request.form.get("ids", "").split(",") if x.strip().isdigit()]; d = db()
    for i in ids: purge_employee(d, "id=?", (i,))
    d.commit(); flash(f"{len(ids)} employee(s) deleted.")
    return redirect(request.form.get("next") if request.form.get("next") in ("/employees", "/attendance") else "/employees")

# ---------- Excel export (.xlsx) ----------
EXPORTS = {
    "employees": "SELECT * FROM employees",
    "attendance": "SELECT e.name, a.day, a.status FROM attendance a JOIN employees e ON e.id=a.emp_id ORDER BY a.day",
    "expenses": "SELECT * FROM expenses",
    "inventory": "SELECT * FROM inventory",
}

@app.route("/exports")
@login_required("admin", "hr")
def exports():
    names = [n for n in EXPORTS if session["role"] == "admin" or n in ("employees", "attendance")]
    return page("""<h2>Export to Excel</h2><div class="card">{% for n in names %}<p><a href="/export/{{n}}.xlsx">Download {{n}}.xlsx</a></p>{% endfor %}</div>""", names=names)

@app.route("/export/<name>.xlsx")
@login_required("admin", "hr")
def export(name):
    if name not in EXPORTS or (session["role"] == "hr" and name not in ("employees", "attendance")):
        return "Not found", 404
    from openpyxl import Workbook
    rows = db().execute(EXPORTS[name]).fetchall()
    wb = Workbook(); ws = wb.active; ws.title = name
    if rows: ws.append(list(rows[0].keys()))
    for r in rows: ws.append(list(r))
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return send_file(buf, download_name=name + ".xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


# =====================================================================
# EDIT / DELETE (employees, expenses, inventory, users)
# =====================================================================
CFG = {
    "employees": dict(roles=("admin", "hr"), back="/employees", fields=EMP_FIELDS),
    "expenses": dict(roles=("admin",), back="/expenses", fields=[("day", "Date", "date"), ("category", "Category", "text"),
        ("description", "Description", "text"), ("amount", "Amount", "number")]),
    "inventory": dict(roles=("admin",), back="/inventory", fields=[("item", "Item", "text"), ("sku", "SKU", "text"),
        ("qty", "Quantity", "number"), ("unit", "Unit", "text"), ("min_qty", "Min qty", "number"), ("price", "Unit price", "number")]),
    "overtime": dict(roles=("admin", "hr"), back="/overtime", fields=[("emp_id", "Employee ID", "number"), ("day", "Date", "date"), ("hours", "Hours", "number")]),
    "loans": dict(roles=("admin", "hr"), back="/loans", fields=[("emp_id", "Employee ID", "number"), ("amount", "Total amount", "number"),
        ("installment", "Monthly installment", "number"), ("start_month", "Start month", "month"), ("note", "Note", "text")]),
    "holidays": dict(roles=("admin", "hr"), back="/holidays", fields=[("day", "Date", "date"), ("name", "Holiday name", "text")]),
    "users": dict(roles=("admin",), back="/users", fields=[("username", "Username", "text"),
        ("role", "Role", "select"), ("emp_id", "Employee ID (for employee role)", "number")]),
}

@app.route("/edit/<table>/<int:i>", methods=["GET", "POST"])
@login_required()
def edit(table, i):
    c = CFG.get(table)
    if not c or session["role"] not in c["roles"]: return "Access denied", 403
    d = db(); row = d.execute(f"SELECT * FROM {table} WHERE id=?", (i,)).fetchone()  # table name is whitelisted via CFG
    if not row: return "Not found", 404
    if request.method == "POST":
        new_id = i
        if table == "employees":   # Changing the employee ID
            try:
                new_id = int(request.form.get("new_id", ""))
                if new_id < 1: raise ValueError
            except ValueError:
                flash("Employee ID must be a whole number of 1 or more."); return redirect(request.path)
            if new_id != i and d.execute("SELECT 1 FROM employees WHERE id=?", (new_id,)).fetchone():
                flash(f"ID {new_id} is already used by another employee."); return redirect(request.path)
            raw = request.form.get("cnic", "").strip(); cn = fmt_cnic(raw)
            if raw and not cn:
                flash("CNIC must contain exactly 13 digits."); return redirect(request.path)
            if cn and d.execute("SELECT 1 FROM employees WHERE cnic=? AND id<>? AND active=1", (cn, i)).fetchone():
                flash("Another employee already has this CNIC."); return redirect(request.path)
        cols, vals = [], []
        for k, _, t in c["fields"]:
            v = request.form.get(k, "").strip()
            if table == "employees" and k == "cnic": v = fmt_cnic(v) or ""
            cols.append(k); vals.append((float(v) if v else None) if t == "number" else v)
        if table == "users":
            if row["username"] == session["user"] and request.form.get("role") != "admin":
                flash("You cannot remove your own admin role."); return redirect(request.path)
            if request.form.get("password"):
                cols.append("password"); vals.append(generate_password_hash(request.form["password"]))
        try:
            d.execute(f"UPDATE {table} SET " + ",".join(k + "=?" for k in cols) + " WHERE id=?", vals + [i])
            if new_id != i:   # move attendance and login to the new ID too
                d.execute("UPDATE attendance SET emp_id=? WHERE emp_id=?", (new_id, i))
                d.execute("UPDATE users SET emp_id=? WHERE emp_id=?", (new_id, i))
                for t in ("leaves", "overtime", "loans"):
                    d.execute(f"UPDATE {t} SET emp_id=? WHERE emp_id=?", (new_id, i))
                d.execute("UPDATE employees SET id=? WHERE id=?", (new_id, i))
                if DATABASE_URL:
                    d.execute("SELECT setval(pg_get_serial_sequence('employees','id'), COALESCE((SELECT MAX(id) FROM employees),0)+1, false)")
            d.commit()
        except Exception:
            d.rollback(); flash("Could not save (this username may already exist)."); return redirect(request.path)
        flash("Changes saved."); return redirect(c["back"])
    depts, desigs = lists(d) if table == "employees" else ([], [])
    return page(FM + """<h2>Edit {{ {'employees':'employee','expenses':'expense','inventory':'inventory item','overtime':'overtime entry','loans':'advance / loan','users':'user'}[table]|capitalize }}</h2>
    <div class="card" style="max-width:{{'1000px' if table=='employees' else '600px'}}"><form method="post"><div class="fgrid">
    {% if table=='employees' %}<div class="fld"><label>Employee ID</label><input name="new_id" type="number" min="1" value="{{row['id']}}" required></div>{% endif %}
    {% for k,label,t in fields %}{{field(k,label,t,row[k])}}{% endfor %}
    {% if table=='users' %}<div class="fld"><label>New password (leave blank to keep the current one)</label><input name="password" type="password"></div>{% endif %}
    </div><p style="margin-top:22px"><button>Save changes</button> <a class="cx" href="{{back}}">Cancel</a></p></form></div>""",
        table=table, fields=c["fields"], row=row, back=c["back"], depts=depts, desigs=desigs)

@app.route("/delete/<table>/<int:i>", methods=["POST"])
@login_required()
def delete(table, i):
    c = CFG.get(table)
    if not c or session["role"] not in c["roles"]: return "Access denied", 403
    d = db()
    if table == "employees":
        purge_employee(d, "id=?", (i,))   # removes the employee, attendance and login
    elif table == "users":
        r = d.execute("SELECT * FROM users WHERE id=?", (i,)).fetchone()
        if r and r["username"] == session["user"]:
            flash("You cannot delete your own account."); return redirect(c["back"])
        if r and r["role"] == "admin" and d.execute("SELECT COUNT(*) FROM users WHERE role='admin'").fetchone()[0] <= 1:
            flash("The last admin cannot be deleted."); return redirect(c["back"])
        d.execute("DELETE FROM users WHERE id=?", (i,))
    else:
        d.execute(f"DELETE FROM {table} WHERE id=?", (i,))
    d.commit(); flash("Deleted."); return redirect(c["back"])

init()  # create the database under gunicorn as well

if __name__ == "__main__":
    # To run online: host="0.0.0.0". In production: gunicorn app:app + HTTPS (Nginx)
    app.run(host="0.0.0.0", port=5000, debug=False)
